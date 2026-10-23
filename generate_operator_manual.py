#!/usr/bin/env python3
"""Generate the 5G+IMS AI remote-assistance operator manual."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "manual_assets"
OUTPUT_DIR = ROOT / "output"
OUTPUT = OUTPUT_DIR / "5G+IMS_AI智能远程协助系统操作员使用说明书_20261023_V2.0.docx"

BLUE = "155A9C"
DARK_BLUE = "103C68"
LIGHT_BLUE = "EAF3FB"
PALE_BLUE = "F4F8FC"
ORANGE = "E67E22"
LIGHT_ORANGE = "FFF3E7"
GRAY = "666666"
LIGHT_GRAY = "F2F2F2"
RED = "C00000"
GREEN = "2E7D32"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, edge_data in edges.items():
        tag = f"w:{edge_name}"
        edge = borders.find(qn(tag))
        if edge is None:
            edge = OxmlElement(tag)
            borders.append(edge)
        for key, value in edge_data.items():
            edge.set(qn(f"w:{key}"), str(value))


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_header_safe(table) -> None:
    if table.rows:
        set_repeat_table_header(table.rows[0])


def set_cell_text(cell, text: str, *, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(9)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[i], BLUE)
        if widths:
            table.rows[0].cells[i].width = Cm(widths[i])
    set_repeat_table_header_safe(table)
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            set_cell_text(cells[i], str(value))
            if widths:
                cells[i].width = Cm(widths[i])
            if row_index % 2:
                set_cell_shading(cells[i], PALE_BLUE)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def set_run_font(run, size=10.5, bold=False, color=None, font="微软雅黑") -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(5)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True, color=DARK_BLUE)
        r = p.add_run(text[len(bold_prefix) :])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.37)
        p.paragraph_format.space_after = Pt(2)
        set_run_font(p.add_run(item))


def add_steps(doc: Document, steps: list[str]) -> None:
    for index, step in enumerate(steps, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        p.paragraph_format.space_after = Pt(5)
        number = p.add_run(f"{index}. ")
        set_run_font(number, bold=True, color=BLUE)
        set_run_font(p.add_run(step))


def add_callout(doc: Document, title: str, text: str, kind: str = "tip") -> None:
    colors = {
        "tip": (LIGHT_BLUE, BLUE),
        "warn": (LIGHT_ORANGE, ORANGE),
        "danger": ("FCE8E6", RED),
        "ok": ("E8F5E9", GREEN),
    }
    fill, accent = colors[kind]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(
        cell,
        left={"val": "single", "sz": "24", "color": accent},
        top={"val": "nil"},
        bottom={"val": "nil"},
        right={"val": "nil"},
    )
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    set_run_font(p.add_run(f"{title}  "), bold=True, color=accent)
    set_run_font(p.add_run(text), size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_picture(
    doc: Document,
    filename: str,
    caption: str,
    *,
    width_cm: float = 16.2,
    required: bool = False,
) -> None:
    path = ASSETS / filename
    if not path.exists():
        if required:
            raise FileNotFoundError(path)
        add_callout(doc, "界面占位", f"截图“{caption}”将在现网采集完成后自动插入。", "tip")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    r = cp.add_run(caption)
    set_run_font(r, size=9, color=GRAY)


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    p = doc.add_paragraph()
    run = p.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "打开文档后按 Ctrl+A，再按 F9 更新目录"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, fld_sep, placeholder, fld_end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(5)

    heading_specs = {
        "Title": (28, DARK_BLUE, "微软雅黑"),
        "Subtitle": (16, GRAY, "微软雅黑"),
        "Heading 1": (18, DARK_BLUE, "微软雅黑"),
        "Heading 2": (14, BLUE, "微软雅黑"),
        "Heading 3": (11.5, DARK_BLUE, "微软雅黑"),
    }
    for style_name, (size, color, font) in heading_specs.items():
        style = doc.styles[style_name]
        style.font.name = font
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if style_name != "Heading 1" else 18)
        style.paragraph_format.space_after = Pt(6)

    caption = doc.styles["Caption"]
    caption.font.name = "微软雅黑"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(GRAY)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    set_run_font(
        hp.add_run("5G+IMS AI智能远程协助系统  |  操作员使用说明书"),
        size=8.5,
        color=GRAY,
    )
    bottom = OxmlElement("w:pBdr")
    border = OxmlElement("w:bottom")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "6")
    border.set(qn("w:space"), "1")
    border.set(qn("w:color"), BLUE)
    bottom.append(border)
    hp._p.get_or_add_pPr().append(bottom)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(fp.add_run("内部使用  ·  第 "), size=8.5, color=GRAY)
    run = fp.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])
    set_run_font(fp.add_run(" 页"), size=8.5, color=GRAY)

    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def cover(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("5G+IMS")
    set_run_font(r, size=22, bold=True, color=BLUE)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("AI智能远程协助系统")
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("操作员使用说明书")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(50)
    set_run_font(p.add_run("Web操作台 · 改版功能适配版"), size=15, color=GRAY)

    line = doc.add_table(rows=1, cols=1)
    line.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = line.cell(0, 0)
    set_cell_shading(cell, BLUE)
    cell.height = Cm(0.08)

    info = [
        ("文档版本", "V2.0"),
        ("适用系统", "5G+IMS AI远程协助系统 Web操作台"),
        ("验证日期", "2026年10月23日"),
        ("文档属性", "内部操作指导文件"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for key, value in info:
        cells = table.add_row().cells
        set_cell_text(cells[0], key, bold=True, color=DARK_BLUE)
        set_cell_text(cells[1], value)
        cells[0].width = Cm(4)
        cells[1].width = Cm(8)
        for c in cells:
            set_cell_border(
                c,
                bottom={"val": "single", "sz": "4", "color": "D9E2F3"},
                top={"val": "nil"},
                left={"val": "nil"},
                right={"val": "nil"},
            )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    set_run_font(p.add_run("未经授权不得对外传播"), size=10, color=RED, bold=True)
    add_page_break(doc)


def front_matter(doc: Document) -> None:
    doc.add_heading("文档说明", level=1)
    add_body(
        doc,
        "本说明书依据2026年10月23日可访问的现网Web操作台逐页核验编制，适用于普通操作员完成语音调度、视频会议、会议会商和录音录像查询等日常工作。文中的菜单名称、字段名称、按钮状态和操作路径均以账号实际可见界面为准；不同权限、不同部署参数或后续升级可能使局部界面有所差异。",
    )
    add_callout(
        doc,
        "重要说明",
        "本手册不记录登录密码，不提供绕过授权、删除保护或审计机制的方法。涉及真实呼叫、会议启动、记录删除等会影响生产业务的动作，仅说明规范流程，编制过程中未在现网执行。",
        "warn",
    )
    doc.add_heading("版本修订记录", level=2)
    add_table(
        doc,
        ["版本", "日期", "修订内容", "状态"],
        [
            [
                "V2.0",
                "2026-10-23",
                "按现网改版界面重新编制；覆盖四个操作模块、会议参数、录音录像和异常处理。",
                "现行",
            ],
            [
                "V1.x",
                "2026-07-27",
                "原操作员手册版本。本次未取得原文件，版式未沿用。",
                "历史",
            ],
        ],
        [2, 3, 10, 2],
    )
    doc.add_heading("适用范围与读者", level=2)
    add_bullets(
        doc,
        [
            "值班调度员：使用语音、视频和会议能力完成日常调度与应急协同。",
            "会议操作员：维护会议室、选择参会成员并在授权后启动会议。",
            "录音录像管理员：按授权查询、播放和清理业务记录。",
            "班组负责人：依据交接清单检查账号状态、设备状态和业务留痕。",
            "系统运维人员：参考故障现象和操作侧检查项协助定位问题。",
        ],
    )
    doc.add_heading("符号约定", level=2)
    add_table(
        doc,
        ["标识", "含义", "处理原则"],
        [
            ["提示", "提高效率或避免误操作的建议", "操作前阅读，按现场情况采用"],
            ["注意", "可能影响通话、会议或数据的事项", "确认条件满足后再继续"],
            ["禁止", "可能造成业务中断、数据丢失或越权的行为", "不得执行"],
            ["现网核验", "已在指定账号界面中实际看到的功能", "菜单名称以截图为准"],
        ],
        [2.5, 6, 6.5],
    )
    doc.add_heading("目录", level=1)
    doc.add_heading("章节导读", level=2)
    add_table(
        doc,
        ["章节", "内容"],
        [
            ["第1章", "系统定位、功能范围和业务链路"],
            ["第2章", "终端、网络、证书及操作前检查"],
            ["第3章", "登录、退出和三类在线状态"],
            ["第4章", "操作台区域、按钮状态和通用规则"],
            ["第5章", "语音调度、拨号、当前通话和最近通话"],
            ["第6章", "视频会议建立、会中控制和结束"],
            ["第7章", "会议室查询、配置、成员和启动"],
            ["第8章", "录音录像查询、播放和删除控制"],
            ["第9章", "AI能力现网边界和结果使用原则"],
            ["第10章", "日常、应急、远程协助和事后核验场景"],
            ["第11章", "故障处理、速查表和上报模板"],
            ["第12章", "账号安全、审计、数据保护和高风险操作"],
            ["第13章", "上岗、下岗和日常维护检查"],
            ["附录", "快速操作卡、状态释义和文档更新"],
        ],
        [3, 13],
    )
    add_callout(
        doc,
        "使用说明",
        "本页按业务主题列出全部章节。可在Word导航窗格中单击标题快速跳转，也可使用Ctrl+F搜索模块或按钮名称。",
        "tip",
    )
    add_page_break(doc)


def chapter_1(doc: Document) -> None:
    doc.add_heading("1 系统概述", level=1)
    doc.add_heading("1.1 系统定位", level=2)
    add_body(
        doc,
        "5G+IMS AI智能远程协助系统面向园区生产调度、现场协同和应急处置场景，将IMS语音、视频会议、会议会商和录音录像统一到一个Web操作台。操作员无需在多个客户端之间频繁切换，即可从通讯录选择人员、发起语音或视频业务、维护固定会议室，并在会后检索录音录像记录。",
    )
    add_body(
        doc,
        "系统标题包含“AI智能远程协助”，但本次账号可见的Web操作台没有独立的“AI助手”“AI降噪”“智能报告”或“安全审计”菜单。此类能力可能由平台后台、媒体链路、移动终端或特定项目授权自动提供。操作员应以实际出现的按钮为准，不得根据产品宣传材料虚构操作入口；如项目启用新的AI控件，应由管理员补充专项操作页并完成现场培训。",
    )
    doc.add_heading("1.2 已核验的功能范围", level=2)
    add_table(
        doc,
        ["一级模块", "主要任务", "本账号可见状态"],
        [
            ["语音调度", "通讯录检索、号码拨叫、语音/视频呼叫、当前通话、最近通话", "可见"],
            ["视频会议", "参会成员选择、会议画面与会中控制", "可见"],
            ["会议会商", "会议室查询、新增、修改、删除、启动和视频窗口", "可见"],
            ["录音录像", "按号码、日期和类型查询；播放及按权限删除", "可见"],
            ["后台管理", "账号、权限、路由、审计和系统参数", "本账号不可见"],
        ],
        [3, 9.5, 3.5],
    )
    doc.add_heading("1.3 业务链路", level=2)
    add_body(
        doc,
        "一次完整业务通常经历“登录与状态确认—选择人员或会议室—建立媒体会话—会中控制—正常结束—录音录像留存—交接记录”七个环节。任何一个环节出现异常，都应先保留页面提示和发生时间，再按本手册的故障检查顺序处理，避免连续重复发起造成对端骚扰或会议资源占用。",
    )
    add_callout(
        doc,
        "现网核验结论",
        "登录后顶部显示“状态服务已连接”、操作员分机状态和在线开关时，才具备稳定执行业务的基础条件。",
        "ok",
    )


def chapter_2(doc: Document) -> None:
    doc.add_heading("2 使用前准备", level=1)
    doc.add_heading("2.1 终端与网络要求", level=2)
    add_bullets(
        doc,
        [
            "建议使用管理员认可的最新版Chrome或Edge浏览器，显示分辨率不低于1920×1080。",
            "操作终端应接入允许访问系统地址的业务网络；代理、防火墙和安全软件不得阻断HTTPS、WebSocket及音视频媒体流。",
            "语音业务需正确连接耳麦或话筒；视频业务需正确连接摄像头，并在浏览器中授予麦克风和摄像头权限。",
            "操作前关闭可能独占摄像头或麦克风的会议软件，确认系统音量、输入设备和输出设备选择正确。",
            "应急值守岗位应配置备用耳麦、备用浏览器和管理员联系方式。",
        ],
    )
    doc.add_heading("2.2 访问地址与证书", level=2)
    add_body(
        doc,
        "当前Web入口为 https://36.141.81.170:8091/webuser/ 。建议由单位门户、受控书签或统一运维页面进入，避免手工输入错误地址。若浏览器提示证书不受信任，必须先向系统管理员核对服务器地址和证书指纹；确认属于本系统后再按单位安全制度处理，不能在来源不明的页面继续输入账号。",
    )
    add_callout(
        doc,
        "安全要求",
        "登录地址、账号和密码应分开保管。本手册仅保留系统地址，不写入生产密码。首次登录后应按管理员要求修改初始密码。",
        "danger",
    )
    doc.add_heading("2.3 操作前检查", level=2)
    add_table(
        doc,
        ["检查项", "正常标准", "异常处理"],
        [
            ["网络", "登录页可在合理时间内加载", "检查白名单、VPN、代理和8091端口"],
            ["浏览器", "页面布局完整，无脚本错误提示", "清理站点缓存后重新登录"],
            ["音频", "耳麦可被系统识别，输入输出正常", "重新选择设备或更换USB接口"],
            ["视频", "摄像头未被其他程序占用", "关闭其他会议软件并重新授权"],
            ["账号", "账号处于有效期且已分配操作权限", "联系管理员核验账号与路由权限"],
            ["值班记录", "已记录班次、操作员和交接事项", "补齐交接后再执行业务"],
        ],
        [3, 7, 6],
    )


def chapter_3(doc: Document) -> None:
    doc.add_heading("3 登录、退出与状态确认", level=1)
    doc.add_heading("3.1 登录系统", level=2)
    add_steps(
        doc,
        [
            "在受控浏览器中打开系统入口，确认地址栏中的服务器IP、端口和路径正确。",
            "在“请输入用户名”框中输入本人操作账号，在“请输入密码”框中输入密码。不得使用他人账号代班。",
            "单击“登录”。系统认证成功后自动进入操作台；当前部署默认进入语音调度页面。",
            "查看页面顶部系统名称、当前账号、状态服务和分机在线状态。只有信息均符合本班次账号时才继续操作。",
        ],
    )
    add_picture(doc, "010-系统登录页.png", "图3-1 系统登录页面")
    doc.add_heading("3.2 登录后的三类状态", level=2)
    add_table(
        doc,
        ["状态位置", "正常显示", "说明"],
        [
            ["页面顶部", "状态服务已连接", "Web操作台与后台状态通道已建立"],
            ["账号旁", "在线空闲", "当前分机已注册且未处于通话中"],
            ["在线开关", "在线", "允许接收和发起受支持的业务"],
        ],
        [4, 4, 8],
    )
    add_body(
        doc,
        "“页面能够打开”不等于业务链路已经就绪。若顶部显示“状态服务未连接”，通讯录可能仍能展示，但呼叫状态、会议成员状态和控制指令可能无法实时更新。遇到这种情况应先刷新页面并等待状态恢复，仍未连接时保存截图和时间点并联系运维。",
    )
    doc.add_heading("3.3 安全退出", level=2)
    add_steps(
        doc,
        [
            "确认当前无进行中的通话或会议，重要录音录像已经形成。",
            "单击页面右上角退出图标，按提示退出当前账号。",
            "确认浏览器重新显示登录页，再关闭浏览器窗口。",
            "公共值班终端交接时，不得仅关闭标签页或让浏览器保存密码。",
        ],
    )


def chapter_4(doc: Document) -> None:
    doc.add_heading("4 操作台界面与通用规则", level=1)
    doc.add_heading("4.1 页面区域", level=2)
    add_table(
        doc,
        ["区域", "内容", "操作要点"],
        [
            ["顶栏", "系统时间、状态服务、系统名称、账号、分机状态、退出", "每次操作前先看状态"],
            ["功能导航", "语音调度、视频会议、会议会商、录音录像", "当前模块以高亮显示"],
            ["业务工作区", "通讯录、会议列表、视频窗口或记录列表", "随模块切换"],
            ["查询区", "号码、备注、日期、类型等条件", "查询后可用“重置”清空"],
            ["操作区", "呼叫、邀请、布局、播放、删除等按钮", "灰色表示当前条件不满足"],
        ],
        [3, 7, 6],
    )
    doc.add_heading("4.2 按钮状态与防误触", level=2)
    add_body(
        doc,
        "系统会根据业务状态控制按钮是否可用。例如未选择联系人时，语音呼叫和视频呼叫按钮保持灰色；未建立会议时，邀请、布局、成员、禁言和结束等会中控制可能不可用。操作员不应反复点击灰色按钮，也不应通过刷新或浏览器开发工具绕过状态限制。",
    )
    add_bullets(
        doc,
        [
            "绿色按钮通常表示发起、查看或播放类动作；红色按钮通常表示结束或删除类高风险动作。",
            "列表复选框用于批量选择。执行删除前要再次核对筛选条件和选中数量。",
            "“重置”只清空当前表单或恢复已加载数据，不代表撤销已经提交的业务。",
            "页面显示加载中时应等待完成，不要连续刷新或重复提交。",
        ],
    )
    doc.add_heading("4.3 日期、列表和分页", level=2)
    add_body(
        doc,
        "会议和录音录像列表可能使用分页显示。查询条件改变后，系统通常回到第一页；若未找到目标记录，应检查日期、号码、类型以及当前页码。列表列宽不足时可横向滚动查看“操作”列，不能因为看不到按钮就判断功能不存在。",
    )


def chapter_5(doc: Document) -> None:
    doc.add_heading("5 语音调度", level=1)
    doc.add_heading("5.1 功能说明", level=2)
    add_body(
        doc,
        "语音调度是操作员的主要工作台。页面左侧提供Web通讯录和部门树，中部提供号码输入及语音/视频呼叫入口，右侧或下方展示快捷联系人、当前通话和最近通话。现网通讯录按部门分组并显示数量，便于从大量分机中快速定位对象。",
    )
    add_picture(doc, "020-语音调度工作台.png", "图5-1 语音调度工作台")
    doc.add_heading("5.2 从通讯录选择联系人", level=2)
    add_steps(
        doc,
        [
            "确认顶部状态服务已连接，分机处于在线空闲状态。",
            "在通讯录区域展开目标部门；需要改变部门显示深度时，使用“部门显示层级”按钮。",
            "在“搜索部门、姓名或号码”中输入关键字。搜索结果可能同时匹配部门、姓名和分机号。",
            "单击目标联系人并核对姓名、号码和所属部门。选择成功后，相应呼叫按钮由灰色变为可用。",
            "按业务需要选择语音呼叫或视频呼叫。接通后在“当前通话”区域观察状态。",
        ],
    )
    doc.add_heading("5.3 直接号码拨叫", level=2)
    add_steps(
        doc,
        [
            "在被叫号码输入框中输入完整分机号或管理员认可的外线号码。",
            "口头复核号码，避免将敏感内容呼叫到错误对象。",
            "单击语音呼叫按钮。若号码不存在，结束尝试并检查号码来源，不要连续重复呼叫。",
        ],
    )
    doc.add_heading("5.4 使用拨号键盘", level=2)
    add_body(
        doc,
        "单击“拨号键盘”可展开数字键、星号和井号。键盘适合临时输入号码或在业务允许时发送DTMF按键。删除末位号码用于修正输入，键盘中的呼叫按钮在号码满足条件后可用。",
    )
    add_picture(doc, "021-语音调度拨号键盘.png", "图5-2 拨号键盘")
    add_callout(
        doc,
        "注意",
        "星号、井号和通话中按键可能触发对端业务菜单。未明确业务要求时不要随意发送。",
        "warn",
    )
    doc.add_heading("5.5 当前通话与最近通话", level=2)
    add_body(
        doc,
        "“当前通话”显示正在进行的会话数量和状态，是确认呼叫是否真正建立的重要依据。结束通话后，记录进入“最近通话”。最近通话区域支持刷新；选中有效记录后可按权限回拨。回拨前仍需确认对象和班次，不能把历史号码视为当前联系人。",
    )
    doc.add_heading("5.6 本地视频预览", level=2)
    add_body(
        doc,
        "发起视频业务前可使用“本地视频预览”检查摄像头取景、光线和隐私背景。预览不等于已经向对端发送画面。若浏览器请求摄像头权限，应选择本岗位指定设备；离开岗位前关闭预览，避免摄像头长期占用。",
    )
    doc.add_heading("5.7 语音调度异常处理", level=2)
    add_table(
        doc,
        ["现象", "优先检查", "建议处理"],
        [
            ["呼叫按钮灰色", "是否选择联系人或填写号码", "重新选择目标并确认在线状态"],
            ["分机不存在", "号码是否正确、分机是否已配置", "停止重复呼叫，联系通讯录管理员"],
            ["呼叫无声音", "耳麦、系统音量、浏览器权限", "切换正确输入输出设备"],
            ["状态服务未连接", "网络和WebSocket通道", "刷新一次；仍异常则联系运维"],
            ["视频无画面", "摄像头权限和设备占用", "关闭其他会议软件后重试"],
            ["最近通话不更新", "状态服务、刷新按钮", "刷新列表并记录发生时间"],
        ],
        [4, 6, 6],
    )


def chapter_6(doc: Document) -> None:
    doc.add_heading("6 视频会议", level=1)
    doc.add_heading("6.1 页面构成", level=2)
    add_body(
        doc,
        "视频会议页面由通讯录、视频窗口和会中控制栏组成。未建立会议时，邀请、布局、成员、禁言和结束等按钮可能保持灰色；“更多”菜单仍可展示可用控制项，但具体设置只有在会议资源和权限满足时才会生效。",
    )
    add_picture(doc, "030-视频会议工作台.png", "图6-1 视频会议工作台")
    doc.add_heading("6.2 建立视频协作", level=2)
    add_steps(
        doc,
        [
            "确认摄像头和麦克风已经连接，浏览器权限正确，分机状态为在线空闲。",
            "在左侧通讯录中展开部门或搜索人员，核对目标成员身份。",
            "根据页面可用状态选择成员并发起视频业务。建立过程中不要重复点击。",
            "会议建立后观察视频窗口和成员状态，确认双方声音和画面正常。",
            "需要扩大会话时使用“邀请”，从通讯录中继续选择成员。",
        ],
    )
    doc.add_heading("6.3 会中控制", level=2)
    add_table(
        doc,
        ["按钮", "作用", "操作注意"],
        [
            ["邀请", "向当前会议增加成员", "邀请前确认对象和会议敏感等级"],
            ["布局", "调整视频画面布局", "人数变化后检查主画面是否合适"],
            ["成员", "查看当前参会成员", "核对异常掉线、重复成员"],
            ["禁言", "执行全体禁言控制", "发布重要指令前说明禁言原因"],
            ["更多", "展开带宽、帧率、位置、主屏、比例和锁定", "按网络和场景调整"],
            ["结束", "结束当前视频会议", "高风险动作，先确认无人继续发言"],
        ],
        [2.5, 6.5, 7],
    )
    doc.add_heading("6.4 更多控制", level=2)
    add_picture(doc, "031-视频会议更多控制.png", "图6-2 视频会议“更多”控制")
    add_bullets(
        doc,
        [
            "带宽：限制或调整会议媒体带宽。网络拥塞时应由值班负责人或运维指导调整。",
            "帧率：控制视频流畅度。降低帧率可减小网络压力，但会降低运动画面连续性。",
            "位置：调整成员画面在布局中的位置，适用于固定主会场和现场画面。",
            "主屏：指定重点显示画面。切换后应确认其他成员仍可见。",
            "比例：调整画面显示比例，避免图像拉伸或重要区域被裁切。",
            "锁定：限制会议配置或成员变化。锁定前要确认成员已经齐备。",
        ],
    )
    doc.add_heading("6.5 结束会议", level=2)
    add_steps(
        doc,
        [
            "口头确认议题结束，提醒参会者保存必要材料。",
            "确认无正在进行的远程指导动作，现场人员已处于安全状态。",
            "单击“结束”，阅读确认提示后再提交。",
            "返回空闲页面后检查录音录像模块是否形成相应记录。",
        ],
    )


def chapter_7(doc: Document) -> None:
    doc.add_heading("7 会议会商", level=1)
    doc.add_heading("7.1 会议列表", level=2)
    add_body(
        doc,
        "会议会商用于维护可重复使用的会议室配置。页面包含“会议列表”和“视频窗口”两个页签。列表可按会议室号码和备注查询，并提供添加、修改、删除、查看详情和开始会议等操作。会议室是长期配置对象，不能把它当作一次性呼叫随意创建或删除。",
    )
    add_picture(doc, "040-会议会商列表.png", "图7-1 会议会商列表")
    doc.add_heading("7.2 查询与查看详情", level=2)
    add_steps(
        doc,
        [
            "在会议室号码或会议室备注中输入完整值或关键字。",
            "单击“查询”，核对结果数量和会议室信息。",
            "单击目标行的“查看详情”图标，依次查看基本设置、视频设置、录音录像和会议成员。",
            "详情查看完成后关闭弹窗。不要在不清楚原配置用途时执行修改。",
        ],
    )
    add_picture(doc, "041-会议详情-01-基本设置.png", "图7-2 会议详情—基本设置", width_cm=14.8)
    add_picture(doc, "041-会议详情-02-视频设置.png", "图7-3 会议详情—视频设置", width_cm=14.8)
    add_picture(doc, "041-会议详情-03-录音录像.png", "图7-4 会议详情—录音录像", width_cm=14.8)
    add_picture(doc, "041-会议详情-04-会议成员.png", "图7-5 会议详情—会议成员（通讯录内容已脱敏）", width_cm=14.8)
    doc.add_heading("7.3 新增会议室", level=2)
    add_body(
        doc,
        "新增会议室由四组参数共同构成。基本设置定义会议身份和提示音，视频设置定义媒体质量，录音录像决定是否留痕，会议成员定义默认参会对象。参数之间相互影响，应在一次配置过程中完整检查，而不能只填写会议室号码。",
    )
    doc.add_heading("7.3.1 基本设置", level=3)
    add_table(
        doc,
        ["字段", "作用", "填写要求"],
        [
            ["会议室号码", "会议资源唯一标识", "界面提示最长12个字符；按编号规范填写"],
            ["会议室备注", "说明用途或班组", "界面提示最长12个字符；名称应清晰可辨"],
            ["密码", "控制会议加入权限", "界面提示最长10个字符；通过安全渠道传递"],
            ["进入时播放", "成员进入时播放提示音", "重要会议建议启用"],
            ["退出时播放", "成员退出时播放提示音", "便于主持人识别异常离会"],
            ["一人时播放", "仅一名成员时播放提示", "提示会议尚未形成多人会话"],
            ["等待时播放", "等待阶段播放提示", "按现场环境和业务习惯启用"],
            ["最大会议时长(秒)", "限制会议资源占用时间", "按业务时长预留合理余量"],
        ],
        [4, 6, 6],
    )
    add_picture(doc, "042-添加会议-01-基本设置.png", "图7-6 添加会议—基本设置", width_cm=14.8)
    doc.add_heading("7.3.2 视频设置", level=3)
    add_table(
        doc,
        ["字段", "说明", "配置建议"],
        [
            ["开启视频", "决定会议是否启用视频媒体", "纯语音会商可关闭"],
            ["视频分辨率", "控制图像清晰度", "兼顾终端能力和网络带宽"],
            ["视频布局", "定义多画面排列方式", "按默认参会人数选择"],
            ["视频质量", "控制编码质量等级", "网络不稳定时避免盲目设高"],
            ["视频帧率", "每秒视频帧数", "运动场景高，静态会商可适当降低"],
            ["视频带宽(MB)", "会议媒体带宽参数", "遵循运维给定范围"],
            ["最大成员数量", "限制会议容量", "不得超过平台许可和资源规划"],
            ["显示成员名称", "在画面上显示成员标识", "便于调度识别，注意隐私"],
        ],
        [4, 6, 6],
    )
    add_picture(doc, "042-添加会议-02-视频设置.png", "图7-7 添加会议—视频设置", width_cm=14.8)
    doc.add_heading("7.3.3 录音录像", level=3)
    add_body(
        doc,
        "“开启录音录像”决定会议是否形成媒体文件。启用前要确认业务有留痕需求、存储容量充足并已履行告知义务。录制文件可能包含生产画面、人员声音和敏感信息，后续查询、播放和删除必须遵守权限与保存期限。",
    )
    add_picture(doc, "042-添加会议-03-录音录像.png", "图7-8 添加会议—录音录像", width_cm=14.8)
    doc.add_heading("7.3.4 会议成员", level=3)
    add_steps(
        doc,
        [
            "在可选成员区域选择通讯录和部门，或使用搜索框定位人员。",
            "勾选需要预置的成员，观察已选成员数量和清单变化。",
            "逐一核对姓名、分机号和部门，移除误选成员。",
            "只有明确需要重新选择时才使用“清空”，避免一次移除全部成员。",
        ],
    )
    add_picture(doc, "042-添加会议-04-会议成员.png", "图7-9 添加会议—会议成员（通讯录内容已脱敏）", width_cm=14.8)
    doc.add_heading("7.3.5 提交前复核", level=3)
    add_callout(
        doc,
        "提交检查",
        "会议室号码唯一、备注清晰、密码合规、最大时长合理、视频参数与网络匹配、录制策略符合制度、成员清单无误。全部确认后方可单击“添加”。",
        "warn",
    )
    doc.add_heading("7.4 修改和删除", level=2)
    add_body(
        doc,
        "修改会议室前应记录原参数和变更原因，确认会议未在使用，并与相关班组沟通生效时间。删除属于高风险操作，会影响后续启动和历史流程引用；只有确认会议室停用、无业务依赖且已获得授权后才能执行。出现确认对话框时应再次核对会议室号码，不可凭列表位置判断。",
    )
    doc.add_heading("7.5 开始会议与视频窗口", level=2)
    add_steps(
        doc,
        [
            "在会议列表中定位目标会议室，查看详情并核对成员、视频和录制策略。",
            "确认预置成员已知悉会议安排，现场作业处于可接受状态。",
            "单击目标行的“开始会议”，按页面提示建立会话。",
            "切换到“视频窗口”，观察成员接入和媒体状态。",
            "会议结束后正常结束会话，并到录音录像模块核验记录。",
        ],
    )
    add_picture(doc, "043-会议会商视频窗口.png", "图7-10 会议会商—视频窗口")


def chapter_8(doc: Document) -> None:
    doc.add_heading("8 录音录像", level=1)
    doc.add_heading("8.1 功能与数据范围", level=2)
    add_body(
        doc,
        "录音录像模块集中展示分机通话和会议会商产生的媒体记录。列表字段包括类型、业务类型、主叫号码、主叫名称、被叫或会议号、对应名称、开始时间、结束时间、录音时长、通话状态、文件名和操作。记录是否生成取决于业务配置、媒体是否真正建立以及存储服务状态。",
    )
    add_picture(doc, "050-录音录像查询页.png", "图8-1 录音录像查询页面（业务记录已脱敏）")
    doc.add_heading("8.2 查询记录", level=2)
    add_steps(
        doc,
        [
            "根据已知条件填写主叫号码或被叫/会议号。条件不确定时不要同时填写过多字段。",
            "选择录音日期。跨日业务应分别查询相关日期。",
            "按需要选择类型，缩小分机录音或会议录音范围。",
            "单击“查询”，核对返回数量、时间范围和业务类型。",
            "需要重新开始时单击“重置”，确认条件已清空后再查询。",
        ],
    )
    add_picture(doc, "051-录音录像类型筛选.png", "图8-2 按类型筛选录音录像")
    add_picture(doc, "052-录音录像日期筛选.png", "图8-3 选择录音日期")
    doc.add_heading("8.3 理解通话状态", level=2)
    add_body(
        doc,
        "“正常通话挂机”表示媒体会话按正常流程结束；“分机不存在”表示被叫号码未能匹配有效分机；“媒体超时”表示信令或会议可能已建立，但媒体链路未在规定时间内保持正常。状态异常并不一定意味着完全没有记录，应结合录音时长、文件名和现场反馈判断。",
    )
    doc.add_heading("8.4 播放录音录像", level=2)
    add_steps(
        doc,
        [
            "先确认自己具有查看该记录的业务授权，并确保周围环境适合播放敏感内容。",
            "在目标记录的“操作”列单击播放图标。没有文件名的记录，播放按钮会保持不可用。",
            "系统加载媒体文件后，在弹窗中使用音频或视频播放器控制播放、暂停和进度。",
            "核验结束后关闭播放器，不得使用未经批准的录屏、抓包或浏览器缓存方式复制文件。",
        ],
    )
    doc.add_heading("8.5 删除记录", level=2)
    add_body(
        doc,
        "单条删除和“删除所选”都会影响业务留痕。删除前必须确认保存期限、案件或事件关联、审计要求和备份策略。批量删除时应先通过查询条件缩小范围，再逐页检查复选框；确认对话框出现后仍需再次核对。普通值班操作员在没有书面授权时不得删除录音录像。",
    )
    add_callout(
        doc,
        "禁止",
        "不得为了释放空间、隐藏操作失误或方便演示而删除生产录音录像。发现容量告警时应提交运维处理。",
        "danger",
    )


def chapter_9(doc: Document) -> None:
    doc.add_heading("9 AI能力与自动化功能边界", level=1)
    doc.add_heading("9.1 操作员可见性说明", level=2)
    add_body(
        doc,
        "本次现网核验的账号导航只有语音调度、视频会议、会议会商和录音录像四项，页面中未出现独立AI智能对话、AI降噪、自动报告或安全审计按钮。因此，本手册不编造这些功能的点击路径，也不把普通录音录像播放描述为自动报告。",
    )
    doc.add_heading("9.2 可能采用自动方式提供的能力", level=2)
    add_body(
        doc,
        "若项目后台已经部署AI降噪、语音识别、智能摘要或审计服务，它们可能在媒体链路或服务端自动运行，操作员无需额外开启。是否实际生效，应通过项目配置清单、终端版本、管理员界面或验收用例确认，不能仅凭系统名称判断。",
    )
    add_table(
        doc,
        ["能力", "本次Web界面核验", "操作员处理原则"],
        [
            ["AI降噪", "无独立开关", "按项目配置使用；异常时保留原始场景信息"],
            ["AI智能对话", "无独立入口", "如终端出现AI入口，按专项培训操作"],
            ["录音录像自动报告", "无报告菜单", "先使用录音录像查询；报告能力另行确认"],
            ["安全审计", "无审计菜单", "按账号实名操作，不代表后台没有审计"],
            ["自动发现/配置/诊断/开通", "操作员页面无入口", "属于网管或终端运维范围"],
        ],
        [4, 4.5, 7.5],
    )
    doc.add_heading("9.3 AI结果使用原则", level=2)
    add_bullets(
        doc,
        [
            "AI转写、摘要和报告属于辅助信息，关键指令仍需人工复核。",
            "噪声环境下不能把“已启用降噪”当作声音必然清晰的保证，应进行通话确认。",
            "涉及安全、停送电、危险作业和应急处置的结论，以授权人员确认和正式记录为准。",
            "发现AI结果与原始录音录像不一致时，应保留原始文件并按异常流程上报。",
        ],
    )


def chapter_10(doc: Document) -> None:
    doc.add_heading("10 典型业务操作", level=1)
    doc.add_heading("10.1 日常语音调度", level=2)
    add_steps(
        doc,
        [
            "交班后登录系统，确认状态服务、分机在线状态和耳麦。",
            "从通讯录搜索目标人员，核对姓名、部门和分机号。",
            "发起语音呼叫，接通后先报身份和事由。",
            "通话中复述关键时间、地点和任务，要求对方确认。",
            "正常结束通话，在最近通话或录音录像中确认留痕。",
        ],
    )
    doc.add_heading("10.2 应急多人会商", level=2)
    add_steps(
        doc,
        [
            "根据事件等级选择已配置的应急会议室，查看详情确认成员。",
            "确认会议室视频、录制和最大时长配置符合本次事件要求。",
            "按授权启动会议，进入视频窗口观察成员接入。",
            "主持人宣布会商规则，必要时使用禁言、主屏和布局控制。",
            "形成决策后逐项复述责任人、任务和完成时限。",
            "结束会议并核验录音录像；将异常掉线或媒体超时写入值班记录。",
        ],
    )
    doc.add_heading("10.3 现场视频远程协助", level=2)
    add_steps(
        doc,
        [
            "确认现场人员已处于安全位置，摄像设备不会暴露无关敏感区域。",
            "通过通讯录选择现场终端并建立视频业务。",
            "使用主屏和位置控制突出现场画面，根据网络情况调整带宽或帧率。",
            "远程专家给出指导时，由现场人员复述并确认，危险动作不得仅凭画面判断。",
            "协助结束后正常结束会议，按制度查询录音录像并完成事件记录。",
        ],
    )
    doc.add_heading("10.4 事后录音录像核验", level=2)
    add_steps(
        doc,
        [
            "取得查询授权，明确事件时间、主叫、被叫或会议号。",
            "按日期和号码查询，核对业务类型、开始时间、结束时间和状态。",
            "在受控环境播放目标文件，记录需要引用的时间点。",
            "若状态为媒体超时或文件缺失，保存查询截图并联系运维。",
            "核验结束后关闭播放器，按制度填写查阅记录。",
        ],
    )


def chapter_11(doc: Document) -> None:
    doc.add_heading("11 故障处理", level=1)
    doc.add_heading("11.1 处理原则", level=2)
    add_body(
        doc,
        "故障处理遵循“先保安全、再保业务、后查原因”的顺序。生产现场存在风险时，立即改用备用通信手段，不应长时间停留在页面排障。任何异常都应记录发生时间、账号、模块、目标号码或会议室、页面提示和网络状态，这些信息比笼统描述“系统不能用”更有助于定位。",
    )
    doc.add_heading("11.2 常见问题速查", level=2)
    add_table(
        doc,
        ["问题", "可能原因", "操作员处理", "升级条件"],
        [
            ["登录页打不开", "防火墙、白名单、端口或网络故障", "检查网络和地址，保留超时信息", "多终端均失败"],
            ["登录失败", "账号、密码、权限或账号状态异常", "重新核对一次，禁止连续猜测", "确认输入正确仍失败"],
            ["状态服务未连接", "WebSocket或后台状态服务异常", "刷新一次并等待", "持续超过值班容忍时间"],
            ["分机离线", "注册失败、网络或终端故障", "切换在线状态并检查设备", "无法恢复"],
            ["呼叫按钮不可用", "未选择对象或状态不满足", "重新选择联系人/输入号码", "联系人在线仍不可用"],
            ["呼叫无声", "设备选择、权限或媒体链路异常", "检查耳麦和浏览器权限", "双方设备正常仍无声"],
            ["视频卡顿", "带宽不足、帧率过高或终端性能不足", "减少画面、适当降低参数", "持续影响指挥"],
            ["会议无法启动", "会议资源、成员或权限问题", "查看配置，不重复启动", "提示资源或服务异常"],
            ["录音没有文件", "媒体未建立、录制未开启或存储异常", "核对文件名、时长和状态", "应录制业务未形成文件"],
            ["播放失败", "文件缺失、网络或浏览器格式问题", "刷新一次，换受控浏览器", "多条文件均失败"],
        ],
        [3, 4.5, 5, 4],
    )
    doc.add_heading("11.3 上报信息模板", level=2)
    add_table(
        doc,
        ["字段", "填写示例"],
        [
            ["发生时间", "2026-10-23 10:30（北京时间）"],
            ["操作账号", "本人操作账号，不填写密码"],
            ["模块", "语音调度 / 视频会议 / 会议会商 / 录音录像"],
            ["对象", "被叫号码或会议室号码"],
            ["现象", "页面提示、按钮状态、是否可复现"],
            ["网络与设备", "状态服务、分机状态、耳麦/摄像头情况"],
            ["已做处理", "刷新、重登、换设备等；写明次数"],
            ["证据", "脱敏截图、值班记录、相关录音录像编号"],
        ],
        [4, 12],
    )


def chapter_12(doc: Document) -> None:
    doc.add_heading("12 安全、审计与数据保护", level=1)
    doc.add_heading("12.1 账号安全", level=2)
    add_bullets(
        doc,
        [
            "实行一人一号，禁止共享账号、代输密码或在群聊中发送密码。",
            "离岗前退出系统；短时离开也应锁定操作终端。",
            "发现异常登录、权限变化或身份信息不符时立即停止操作并上报。",
            "浏览器不得保存生产密码，不在非受控终端登录。",
        ],
    )
    doc.add_heading("12.2 音视频与个人信息保护", level=2)
    add_body(
        doc,
        "通讯录、通话记录、会议成员和录音录像可能包含姓名、分机、声音、图像及生产现场信息。查询、播放、截图和导出都应遵循最小必要原则。制作培训材料时应对真实姓名、号码、会议主题和记录内容进行脱敏；本手册中的录音录像列表和会议成员截图已进行遮挡处理。",
    )
    doc.add_heading("12.3 操作留痕", level=2)
    add_body(
        doc,
        "操作员页面没有独立安全审计菜单，不代表操作不被记录。登录、会议配置、呼叫、播放和删除等行为可能由后台日志留痕。所有操作都应能够说明业务目的，禁止尝试规避日志、篡改时间或删除证据。",
    )
    doc.add_heading("12.4 高风险操作清单", level=2)
    add_table(
        doc,
        ["操作", "主要风险", "最低控制要求"],
        [
            ["删除会议室", "固定会商入口失效", "审批、停用确认、记录原配置"],
            ["结束会议", "中断现场协同", "主持人确认、现场安全确认"],
            ["全体禁言", "成员无法及时报告", "说明原因并保留应急发言通道"],
            ["删除录音录像", "业务证据不可恢复", "保存期限核对、授权和审计"],
            ["调整带宽/帧率", "画质或网络负荷变化", "在运维建议范围内调整"],
            ["清空会议成员", "预置成员全部移除", "修改前导出或记录原名单"],
        ],
        [4, 6, 6],
    )


def chapter_13(doc: Document) -> None:
    doc.add_heading("13 班次交接与日常检查", level=1)
    doc.add_heading("13.1 上岗检查清单", level=2)
    add_table(
        doc,
        ["序号", "检查内容", "结果"],
        [
            ["1", "登录账号与当前值班人员一致", "□正常  □异常"],
            ["2", "状态服务已连接，分机在线空闲", "□正常  □异常"],
            ["3", "耳麦、麦克风和摄像头工作正常", "□正常  □异常"],
            ["4", "通讯录能够加载和搜索", "□正常  □异常"],
            ["5", "常用应急会议室配置存在", "□正常  □异常"],
            ["6", "录音录像查询服务可用", "□正常  □异常"],
            ["7", "备用通信方式和运维联系方式明确", "□正常  □异常"],
        ],
        [2, 10, 4],
    )
    doc.add_heading("13.2 下岗交接清单", level=2)
    add_table(
        doc,
        ["序号", "交接内容", "说明"],
        [
            ["1", "未结束的通话、会议和待办事项", "写明对象、状态和下一步"],
            ["2", "异常记录", "写明时间、模块、提示和工单编号"],
            ["3", "重要录音录像", "记录查询条件或文件名，不复制文件"],
            ["4", "会议室变更", "写明变更人、原因和生效时间"],
            ["5", "设备状态", "耳麦、摄像头、显示器及备用设备"],
            ["6", "安全退出", "确认账号已退出且终端已锁定"],
        ],
        [2, 7, 7],
    )
    doc.add_heading("13.3 日常维护建议", level=2)
    add_bullets(
        doc,
        [
            "每班至少检查一次关键会议室和录音录像查询能力。",
            "通讯录变更应由授权管理员维护，操作员发现错误时提交修正，不自行绕过。",
            "定期演练备用通信方式，但演练必须与生产号码和真实应急会议区分。",
            "浏览器、终端和媒体设备升级后应进行语音、视频、会议和录制回归测试。",
        ],
    )


def appendix(doc: Document) -> None:
    doc.add_heading("附录A 快速操作卡", level=1)
    add_table(
        doc,
        ["任务", "最短路径", "关键确认"],
        [
            ["语音呼叫", "语音调度→搜索/输入号码→语音呼叫", "状态服务、对象、号码"],
            ["视频呼叫", "语音调度→选择联系人→视频呼叫", "摄像头、权限、现场安全"],
            ["视频会议", "视频会议→选择成员→建立会议", "成员、声音、画面"],
            ["启动会商", "会议会商→查看详情→开始会议→视频窗口", "成员、录制、最大时长"],
            ["查询录音", "录音录像→日期/号码/类型→查询→播放", "授权、时间、业务类型"],
            ["故障上报", "截图→记录时间和对象→启用备用手段→联系运维", "不反复提交"],
        ],
        [3.5, 8, 4.5],
    )
    doc.add_heading("附录B 页面提示与含义", level=1)
    add_table(
        doc,
        ["提示或状态", "含义", "建议"],
        [
            ["状态服务已连接", "实时状态通道正常", "可继续检查分机状态"],
            ["状态服务未连接", "实时状态通道异常", "刷新一次，仍异常则上报"],
            ["在线空闲", "分机在线且未占用", "可发起业务"],
            ["分机注册成功", "当前分机注册完成", "继续确认音视频设备"],
            ["分机不存在", "目标号码未配置或无效", "核对号码，不重复拨叫"],
            ["媒体超时", "媒体链路未正常维持", "检查网络、终端和会议媒体"],
            ["正常通话挂机", "会话按正常流程结束", "按需核验录音录像"],
        ],
        [4, 6, 6],
    )
    doc.add_heading("附录C 功能核验说明", level=1)
    add_body(
        doc,
        "本手册在不影响生产业务的前提下完成只读核验：实际登录、遍历四个导航模块、查看会议详情、打开新增会议表单、展开拨号键盘和视频更多控制、查询录音录像页面。未执行真实呼叫、会议启动、配置提交、删除或批量删除。相关操作步骤依据实际按钮、字段和状态逻辑编写，并明确标注高风险条件。",
    )
    add_body(
        doc,
        "后续系统升级若新增AI智能对话、AI降噪开关、自动报告、安全审计或终端自动运维入口，应更新本手册的功能范围、页面截图、权限说明和验收记录，并提升文档版本号。",
    )
    doc.add_heading("附录D 文档使用与更新", level=1)
    add_bullets(
        doc,
        [
            "打开Word后按Ctrl+A，再按F9，更新页码和文档字段。",
            "打印时使用A4纸、双面打印，并检查截图是否清晰。",
            "系统菜单、字段或权限变化后，由文档责任人重新核验现网并更新版本记录；生产密码、真实音视频内容和未脱敏个人信息不得写入手册。",
        ],
    )


def build() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    cover(doc)
    front_matter(doc)
    for chapter in (
        chapter_1,
        chapter_2,
        chapter_3,
        chapter_4,
        chapter_5,
        chapter_6,
        chapter_7,
        chapter_8,
        chapter_9,
        chapter_10,
        chapter_11,
        chapter_12,
        chapter_13,
        appendix,
    ):
        chapter(doc)
    doc.core_properties.title = "5G+IMS AI智能远程协助系统操作员使用说明书"
    doc.core_properties.subject = "Web操作台改版功能操作指导"
    doc.core_properties.author = "项目文档组"
    doc.core_properties.keywords = (
        "5G, IMS, AI, 远程协助, 语音调度, 视频会议, 会议会商, 录音录像"
    )
    doc.core_properties.comments = "依据2026-10-23现网操作台核验生成"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
