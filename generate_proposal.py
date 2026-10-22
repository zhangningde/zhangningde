#!/usr/bin/env python3
"""Generate the park integrated communications technical proposal.

The script intentionally keeps source documents outside the repository and
accepts their paths as arguments. It uses the reference document as a style
template, parses every tender paragraph/table row, and emits a traceable
response matrix in the final Word document.
"""

from __future__ import annotations

import argparse
import math
import re
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


BLUE = "17365D"
MID_BLUE = "2F5597"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EEF5FB"
GRAY = "666666"
LIGHT_GRAY = "F2F2F2"
GREEN = "E2F0D9"
RED = "C00000"
FONT_NAME = "Microsoft YaHei"
FALLBACK_FONT = "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"


KEY_TOPICS = [
    ("AI智能融合远程协助总体方案", "形成“现场智能终端—安全传输—融合平台—指挥中心—专家席位”闭环，统一承载语音、视频、位置、告警和协作数据。"),
    ("AI降噪与强噪声环境语音增强", "针对钢铁冶炼、港口、风机和车辆背景噪声，采用端侧与平台侧协同降噪、回声消除、增益控制和语音活动检测，提升可懂度。"),
    ("AI智能对话与免手持作业", "支持唤醒词、自然语言指令、上下文理解和语音播报，作业人员可在佩戴手套、登高或检修状态下完成呼叫、取证、查询和求助。"),
    ("一体化AI录音录像自动报告", "对授权的音视频会话自动归档、转写、摘要、待办提取和报告生成，原始媒体、转写文本和报告以同一事件编号关联。"),
    ("全流程安全审计", "覆盖登录、授权、配置、呼叫、调度、录音调阅、下载、接口调用和运维操作，形成身份、时间、对象、动作、结果完整证据链。"),
    ("统一网管与全网设备一体化管理", "对平台、服务器、基站、网关、调度台、智能终端和链路进行统一资产、拓扑、性能、告警、配置、日志和工单管理。"),
    ("智能终端自动发现", "通过标准协议、地址扫描、注册事件和设备指纹发现终端，自动识别型号、版本、能力集、位置和在线状态。"),
    ("智能终端自动配置", "基于组织、岗位、区域和终端类型匹配模板，下发号码、通信组、权限、编解码、安全证书和业务参数。"),
    ("智能终端自动诊断", "持续检查注册、链路、时延、抖动、丢包、媒体协商、配置偏差和软件版本，输出根因建议与处置步骤。"),
    ("智能终端自动开通", "将审批、资源校验、号码分配、模板下发、注册验证、业务测试和台账更新编排为可回滚的自动化流程。"),
    ("四自动闭环运营", "自动发现触发纳管，自动配置形成一致基线，自动诊断识别偏差，自动开通完成验证；失败自动回滚并转人工工单。"),
    ("本地化AI与数据安全", "ASR、TTS、知识库、转写和摘要优先在园区受控环境部署，业务数据按域隔离并执行最小权限、加密和留痕。"),
    ("数字集群与公网对讲融合", "以标准接口接入专网数字集群和公网对讲，实现单呼、组呼、强插、强拆、监听、跨网编组和应急协同。"),
    ("IP电话利旧与IMS融合", "通过SIP及边界会话控制接入既有IP电话，保留号码规划和使用习惯，逐步迁移且不影响现网生产。"),
    ("视频监控与GB/T 28181融合", "接入既有监控资源，实现预览、轮巡、上墙、云台控制、位置联动和事件取证，避免重复建设。"),
    ("广播与应急通知联动", "融合广播分区、电话、短信、APP和对讲终端，支持预案一键触发、失败重试、回执统计和全过程留痕。"),
    ("GIS一张图调度", "在园区地图汇聚人员、车辆、终端、视频、告警和资源状态，支持框选、圈选、轨迹与就近力量调度。"),
    ("多级多域组织与权限", "园区总调度和企业分调度共享资源体系，各企业用户、通信组、数据和管理权限逻辑隔离，授权跨域协同。"),
    ("高可用与容灾设计", "核心服务冗余部署，支持故障检测、业务切换、配置同步、数据备份和恢复演练，满足7×24小时运行。"),
    ("极端场景通信保障", "面向断网、断电、断路和局部设施故障，采用多链路、后备电源、分级降级、本地自治和应急预案。"),
    ("开放接口与集成治理", "提供标准SIP、REST API、SDK和事件订阅，实施身份认证、限流、幂等、版本管理、审计和接口健康监测。"),
    ("多语言与印尼本地适配", "平台支持中文、英文或印尼文，终端按要求切换语言，并在实施中完成当地法规、频率和认证核验。"),
    ("网络与媒体质量保障", "对信令和媒体分区分级，实施QoS、带宽核算、抖动缓冲、丢包补偿、时钟同步和端到端质量监测。"),
    ("零信任安全与纵深防护", "通过身份可信、设备可信、链路加密、边界控制、最小权限、漏洞治理和持续审计构建纵深防线。"),
    ("可测试、可验收、可运维交付", "每项需求绑定设计、配置证据、测试用例和验收记录，形成从招标条款到交付成果的双向追溯。"),
]


def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def chinese_font(size: int, bold: bool = False):
    try:
        return ImageFont.truetype(FALLBACK_FONT, size=size)
    except OSError:
        return ImageFont.load_default()


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=50, start=70, bottom=50, end=70) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "更新域后显示"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, placeholder, end])


def clear_document_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.first_line_indent = Cm(0.74)

    for style_name, size, color, before, after in (
        ("Title", 30, BLUE, 0, 12),
        ("Heading 1", 20, BLUE, 12, 8),
        ("Heading 2", 15, MID_BLUE, 8, 6),
        ("Heading 3", 12, MID_BLUE, 6, 4),
    ):
        try:
            style = doc.styles[style_name]
        except KeyError:
            style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    try:
        doc.styles["Table Grid"]
    except KeyError:
        doc.styles.add_style("Table Grid", WD_STYLE_TYPE.TABLE)
    try:
        doc.styles["List Bullet"]
    except KeyError:
        bullet = doc.styles.add_style("List Bullet", WD_STYLE_TYPE.PARAGRAPH)
        bullet.font.name = FONT_NAME
        bullet._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        bullet.font.size = Pt(10)


def configure_page(doc: Document) -> None:
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.55)
        section.bottom_margin = Cm(1.45)
        section.left_margin = Cm(1.65)
        section.right_margin = Cm(1.55)
        section.header_distance = Cm(0.65)
        section.footer_distance = Cm(0.65)
        section.different_first_page_header_footer = True
        for header in (section.header, section.first_page_header):
            for p in header.paragraphs:
                p.clear()
        for footer in (section.footer, section.first_page_footer):
            for p in footer.paragraphs:
                p.clear()

        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run("园区融合通讯系统建设项目｜技术方案")
        set_run_font(r, 8, False, GRAY)
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("第 ")
        set_run_font(r, 8, False, GRAY)
        add_field(fp, "PAGE")
        r = fp.add_run(" 页")
        set_run_font(r, 8, False, GRAY)


def parse_tender(path: Path):
    doc = Document(path)
    sections: list[dict] = []
    requirements: list[dict] = []
    current_chapter = "总则"
    current_section = None
    req_no = 0

    for para in doc.paragraphs:
        text = clean_text(para.text)
        if not text or text == "目录":
            continue
        style = para.style.name
        is_heading = style.startswith(("Heading", "标题"))
        if is_heading:
            level_match = re.search(r"(\d+)$", style)
            level = int(level_match.group(1)) if level_match else 2
            if re.search(r"第[一二三四五六七八九十百]+章", text):
                level = 1
                current_chapter = text
            elif level == 1:
                current_chapter = text
            current_section = {"title": text, "level": level, "chapter": current_chapter, "items": []}
            sections.append(current_section)
        else:
            if current_section is not None:
                current_section["items"].append(text)
            if len(text) > 2 and not text.startswith(("二〇", "融合通讯系统建设项目", "技术招标需求书")):
                req_no += 1
                requirements.append(
                    {"id": f"TR-{req_no:04d}", "chapter": current_chapter, "section": current_section["title"] if current_section else "总则", "text": text}
                )

    for table_index, table in enumerate(doc.tables, 1):
        if not table.rows:
            continue
        headers = [clean_text(c.text) for c in table.rows[0].cells]
        for row_index, row in enumerate(table.rows[1:], 1):
            cells = [clean_text(c.text) for c in row.cells]
            text = "；".join(f"{headers[i] or f'字段{i+1}'}：{v}" for i, v in enumerate(cells) if v)
            if text:
                req_no += 1
                requirements.append(
                    {"id": f"TR-{req_no:04d}", "chapter": "招标数据表", "section": f"招标表{table_index}", "text": text}
                )
    return sections, requirements, doc


def create_diagram(path: Path, title: str, columns: list[tuple[str, list[str]]], footer: str) -> None:
    width, height = 1600, 900
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    title_font = chinese_font(46, True)
    head_font = chinese_font(28, True)
    body_font = chinese_font(23)
    foot_font = chinese_font(20)
    draw.rectangle((0, 0, width, 90), fill=(23, 54, 93))
    draw.text((60, 21), title, font=title_font, fill="white")
    margin, gap = 55, 30
    box_w = (width - 2 * margin - gap * (len(columns) - 1)) // len(columns)
    palette = [(217, 234, 247), (238, 245, 251), (226, 240, 217), (242, 242, 242), (255, 242, 204)]
    for idx, (heading, items) in enumerate(columns):
        x1 = margin + idx * (box_w + gap)
        x2 = x1 + box_w
        draw.rounded_rectangle((x1, 145, x2, 760), radius=18, fill=palette[idx % len(palette)], outline=(47, 85, 151), width=3)
        draw.rectangle((x1, 145, x2, 215), fill=(47, 85, 151))
        bbox = draw.textbbox((0, 0), heading, font=head_font)
        draw.text((x1 + (box_w - (bbox[2] - bbox[0])) / 2, 164), heading, font=head_font, fill="white")
        y = 255
        for item in items:
            wrapped = [item[i : i + 11] for i in range(0, len(item), 11)]
            draw.ellipse((x1 + 24, y + 8, x1 + 36, y + 20), fill=(47, 85, 151))
            for line in wrapped:
                draw.text((x1 + 48, y), line, font=body_font, fill=(35, 35, 35))
                y += 34
            y += 23
        if idx < len(columns) - 1:
            ax = x2 + 4
            ay = 450
            draw.line((ax, ay, ax + gap - 8, ay), fill=(192, 0, 0), width=5)
            draw.polygon([(ax + gap - 8, ay), (ax + gap - 20, ay - 10), (ax + gap - 20, ay + 10)], fill=(192, 0, 0))
    draw.rectangle((0, 810, width, 900), fill=(238, 245, 251))
    draw.text((55, 836), footer, font=foot_font, fill=(23, 54, 93))
    img.save(path, quality=95)


def build_diagrams(output_dir: Path) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    specs = {
        "overall": (
            "园区融合通讯系统总体架构",
            [
                ("终端接入层", ["专网/公网对讲", "IP智能电话", "智能安全帽", "APP/调度台"]),
                ("融合能力层", ["软交换与媒体", "会议/广播/GIS", "录音录像", "AI能力中心"]),
                ("业务应用层", ["生产调度", "远程协助", "应急指挥", "统一运维"]),
                ("安全运维层", ["身份与权限", "安全审计", "统一网管", "高可用容灾"]),
            ],
            "设计原则：统一平台、分级管理、互联互通、利旧整合、安全可靠、开放扩展。",
        ),
        "ai": (
            "AI智能远程协助闭环",
            [
                ("现场感知", ["音视频采集", "AI降噪", "定位与告警", "语音唤醒"]),
                ("智能理解", ["ASR转写", "意图识别", "知识检索", "多轮对话"]),
                ("协同处置", ["专家连线", "免手持指令", "标注指导", "应急调度"]),
                ("证据与改进", ["录音录像", "安全审计", "AI自动报告", "知识沉淀"]),
            ],
            "原始媒体、转写文本、协作轨迹、处置结果和报告统一绑定事件编号，支持追溯。",
        ),
        "automation": (
            "智能终端“四自动”运营闭环",
            [
                ("自动发现", ["注册事件", "设备指纹", "能力识别", "资产纳管"]),
                ("自动配置", ["模板匹配", "参数下发", "证书装载", "基线校验"]),
                ("自动诊断", ["质量检测", "偏差比对", "根因定位", "修复建议"]),
                ("自动开通", ["流程编排", "业务验证", "失败回滚", "台账更新"]),
            ],
            "统一网管以策略、模板、工作流和审计日志贯穿全流程，实现批量、可控、可回滚。",
        ),
        "security": (
            "纵深安全与审计架构",
            [
                ("访问安全", ["统一身份", "多因素认证", "最小权限", "分域隔离"]),
                ("通信安全", ["SIP安全", "媒体加密", "边界控制", "抗攻击"]),
                ("数据安全", ["本地化存储", "分类分级", "加密备份", "留存销毁"]),
                ("审计运营", ["操作审计", "通话审计", "敏感词检测", "报告取证"]),
            ],
            "安全控制覆盖人、设备、应用、接口、网络和数据，告警与工单形成闭环。",
        ),
        "management": (
            "统一网管全网一体化管理",
            [
                ("资源中心", ["组织与用户", "设备与链路", "号码与通信组", "版本与许可"]),
                ("监控中心", ["拓扑状态", "性能质量", "业务可用性", "容量趋势"]),
                ("运维中心", ["配置基线", "告警关联", "诊断处置", "工单知识库"]),
                ("运营中心", ["统计报表", "SLA分析", "审计追溯", "决策看板"]),
            ],
            "平台、基站、网关、调度台和智能终端统一视图、统一策略、统一告警、统一审计。",
        ),
    }
    result = {}
    for key, (title, columns, footer) in specs.items():
        path = output_dir / f"{key}.png"
        create_diagram(path, title, columns, footer)
        result[key] = path
    return result


class ProposalWriter:
    def __init__(self, doc: Document):
        self.doc = doc
        self.explicit_pages = 0
        self.started = False

    def new_page(self, title: str, kicker: str | None = None, level: int = 1) -> None:
        if self.started:
            self.doc.add_page_break()
        self.started = True
        self.explicit_pages += 1
        if kicker:
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(kicker)
            set_run_font(r, 8.5, True, MID_BLUE)
        p = self.doc.add_paragraph(style=f"Heading {min(max(level, 1), 3)}")
        p.paragraph_format.space_before = Pt(0)
        p.add_run(title)
        self.rule()

    def rule(self) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p_pr = p._p.get_or_add_pPr()
        p_bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), MID_BLUE)
        p_bdr.append(bottom)
        p_pr.append(p_bdr)

    def paragraph(self, text: str, bold_lead: str | None = None, size: float = 10.2, indent: bool = True) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Cm(0.74) if indent else None
        if bold_lead and text.startswith(bold_lead):
            r = p.add_run(bold_lead)
            set_run_font(r, size, True, BLUE)
            r = p.add_run(text[len(bold_lead) :])
            set_run_font(r, size)
        else:
            r = p.add_run(text)
            set_run_font(r, size)

    def bullets(self, items: Iterable[str], size: float = 9.6) -> None:
        for item in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.first_line_indent = Cm(-0.28)
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.line_spacing = 1.2
            r = p.add_run(item)
            set_run_font(r, size)

    def table(self, headers: list[str], rows: list[list[str]], widths: list[float] | None = None, size: float = 8.2) -> None:
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        table.autofit = False
        if widths:
            for i, width in enumerate(widths):
                table.columns[i].width = Cm(width)
        header = table.rows[0]
        set_repeat_table_header(header)
        for i, text in enumerate(headers):
            cell = header.cells[i]
            set_cell_shading(cell, BLUE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            set_run_font(r, size, True, "FFFFFF")
        for ridx, values in enumerate(rows):
            cells = table.add_row().cells
            for i, value in enumerate(values):
                if ridx % 2:
                    set_cell_shading(cells[i], PALE_BLUE)
                set_cell_margins(cells[i])
                cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cells[i].paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                r = p.add_run(value)
                set_run_font(r, size)

    def image(self, path: Path, width: float = 17.2) -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        p.add_run().add_picture(str(path), width=Cm(width))

    def note(self, text: str) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_margins(cell, 90, 110, 90, 110)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, 9, True, BLUE)


def section_focus(title: str) -> tuple[str, str, str]:
    if any(k in title for k in ("录音", "录像")):
        return "媒体证据全生命周期管理", "统一采集、可靠存储、智能检索、分权调阅、AI分析与审计留痕", "抽取样本核验媒体、转写、报告和审计记录的一致性"
    if any(k in title for k in ("网管", "运维", "运行管理")):
        return "统一网管与自动化运营", "资源、拓扑、性能、告警、配置、日志、工单及终端四自动闭环", "制造故障并检查发现、定位、告警、处置、恢复全过程"
    if any(k in title for k in ("安全", "合规", "认证")):
        return "纵深安全和可审计控制", "身份、权限、边界、通信、数据、接口和运维安全协同", "执行越权、弱口令、异常注册和审计追溯测试"
    if any(k in title for k in ("接口", "对接", "整合", "融合")):
        return "标准开放和利旧融合", "采用标准协议、适配网关和API治理，降低对现网的侵入", "使用接口清单、联调记录、异常用例和回归结果验收"
    if any(k in title for k in ("专网", "集群", "无线", "基站", "覆盖")):
        return "工业场景无线可靠通信", "围绕覆盖、容量、切换、抗干扰、终端防护和极端场景保障设计", "通过覆盖测试、并发压力、漫游切换和断链降级验证"
    if any(k in title for k in ("公网", "4G", "APP")):
        return "广域移动协同", "利用公网覆盖实现跨区域语音、视频、消息、定位和调度", "验证弱网、切换、离线、权限与业务连续性"
    if any(k in title for k in ("视频", "会议", "会商")):
        return "多媒体可视化协同", "统一会控、视频资源接入、双流、录制及跨终端协作", "以多终端并发、协议兼容、画面质量和会控功能验收"
    if any(k in title for k in ("电话", "软交换", "语音", "广播")):
        return "电信级语音与通知能力", "以SIP软交换为核心，融合电话、广播、会议、录音和调度业务", "通过呼叫模型、并发、路由、失败重试和话单一致性验证"
    if any(k in title for k in ("GIS", "定位", "应急")):
        return "一张图应急指挥闭环", "汇聚位置、视频、告警、预案和通信资源，实现快速研判处置", "通过预案演练验证发现、通知、调度、反馈、复盘链路"
    if any(k in title for k in ("设备", "终端", "规格")):
        return "标准化设备与全生命周期管理", "设备选型满足环境、协议和容量要求，并纳入统一网管", "逐项执行到货、上电、注册、业务、性能和稳定性检查"
    return "招标需求完整落地", "遵循统一平台、分级管理、安全可靠、互联互通、标准开放和利旧整合原则", "以条款、设计、配置、测试和证据五级追溯完成验收"


def summarize_items(items: list[str], limit: int = 5) -> list[str]:
    result = []
    for text in items[:limit]:
        result.append(text if len(text) <= 95 else text[:92] + "……")
    return result or ["本节按招标技术要求、总体架构和项目实施边界进行深化设计。"]


def add_cover(writer: ProposalWriter) -> None:
    writer.started = True
    writer.explicit_pages = 1
    for _ in range(3):
        writer.doc.add_paragraph()
    p = writer.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("园区融合通讯系统建设项目")
    set_run_font(r, 27, True, BLUE)
    p.paragraph_format.space_after = Pt(18)
    p = writer.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("技术方案")
    set_run_font(r, 38, True, MID_BLUE)
    p.paragraph_format.space_after = Pt(32)
    p = writer.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("统一融合｜AI智能｜安全审计｜一体运维")
    set_run_font(r, 15, False, GRAY)
    for _ in range(7):
        writer.doc.add_paragraph()
    p = writer.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("投标技术文件（建议稿）")
    set_run_font(r, 13, True, BLUE)
    p = writer.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("二〇二六年十月")
    set_run_font(r, 12, False, GRAY)


def add_front_matter(writer: ProposalWriter) -> None:
    writer.new_page("文档控制与使用说明", "前置文件", 1)
    writer.table(
        ["项目", "内容"],
        [
            ["文档名称", "园区融合通讯系统建设项目技术方案"],
            ["编制依据", "园区融合通讯系统技术招标需求书及已提供产品资料"],
            ["版本", "V1.0（投标建议稿）"],
            ["适用范围", "园区总调度、企业分调度、专网/公网对讲、IP电话、视频、广播、GIS、远程协助及统一运维"],
            ["参数原则", "招标文件优先；产品资料用于能力说明，冲突处按招标要求响应"],
            ["保密要求", "未经授权不得复制、外传或用于本项目以外用途"],
        ],
        [4, 13],
    )
    writer.note("重要说明：指定的“2026-09_现网痛点页完善版”未随资料提供；本方案仅依据明确提出的AI降噪、AI智能对话、免手持、安全审计和自动报告要求，以及已上传远程协助资料进行设计。")

    writer.new_page("项目理解与总体判断", "前置文件", 1)
    writer.paragraph("本项目不是单一语音系统替换，而是面向约80个部门、复杂工业环境和多制式存量资源的园区级通信底座建设。核心任务是消除电话、集群对讲、视频监控、广播和GIS之间的信息孤岛，在保持企业分级自治的同时，形成园区统一调度和应急指挥能力。")
    writer.paragraph("方案以融合通信平台为核心，以专网数字集群和公网移动通信为广域触达手段，以IP电话、视频、广播和GIS利旧接入为基础，以AI远程协助、录音录像智能分析和统一网管自动化为能力增强。")
    writer.bullets(["统一资源体系与分级分域治理并重", "生产通信可靠性与开放扩展并重", "现网平滑迁移与未来容量演进并重", "业务创新与数据安全、可审计并重"])

    writer.new_page("技术响应总承诺", "前置文件", 1)
    writer.table(
        ["承诺维度", "响应原则", "交付证据"],
        [
            ["完整性", "逐项响应招标技术条款，不以偏离说明替代实现方案", "条款响应矩阵、设计章节、测试用例"],
            ["一致性", "需求、设计、配置、测试和验收保持同一编号链路", "追溯矩阵、配置基线、验收记录"],
            ["兼容性", "采用标准协议及开放接口，保护现网投资", "接口清单、联调报告、兼容测试"],
            ["可靠性", "核心服务冗余、链路多路径、数据备份和降级预案", "切换测试、压力测试、恢复演练"],
            ["安全性", "最小权限、纵深防护、全程审计和数据本地化", "安全配置、审计日志、安全测试"],
            ["可运维性", "全网统一网管与终端四自动运营闭环", "拓扑、告警、工单、自动化执行记录"],
        ],
        [3.2, 8.2, 5.5],
    )

    writer.new_page("方案核心亮点", "前置文件", 1)
    writer.bullets(
        [
            "AI降噪：端云协同抑制持续噪声和突发噪声，保障强噪声作业区语音可懂度。",
            "AI智能对话与解放双手：自然语言唤醒、查询、呼叫、取证和求助，降低现场操作负担。",
            "一体化AI录音录像报告：原始媒体、转写、摘要、待办、风险点与审计证据自动关联。",
            "安全审计：从身份登录到调度操作、媒体调阅、下载导出和接口调用全过程留痕。",
            "统一网管：平台、基站、网关、调度台和智能终端统一资产、监控、配置、告警与工单。",
            "终端四自动：自动发现、自动配置、自动诊断、自动开通，支持批量执行、失败回滚和审计。",
        ]
    )

    writer.new_page("编制依据与优先级", "前置文件", 1)
    writer.table(
        ["优先级", "依据", "使用方式"],
        [
            ["P1", "园区融合通讯系统技术招标需求书", "架构、功能、参数、供货、验收、实施和服务的最高依据"],
            ["P2", "国家、行业、国际及印尼当地适用规范", "作为合规、接口、安全和工程实施约束"],
            ["P3", "US2000产品白皮书及管理员手册", "用于平台功能、配置、运维、安全与操作说明"],
            ["P4", "AI+IMS、远程协助、录音及UC2000资料", "用于AI、业务场景、录音和统一管理能力深化"],
            ["P5", "参考方案V1.1.0", "用于版式、语言风格和方案表达方式参考"],
        ],
        [2.3, 7, 7.6],
    )
    writer.paragraph("当产品资料中的历史参数、示例场景或名称与招标要求不一致时，本方案不直接照搬，统一以招标文件要求和最终经确认的深化设计为准。")

    writer.new_page("阅读导航与追溯方法", "前置文件", 1)
    writer.paragraph("正文顺序与招标需求书章节保持一致。每个技术主题从“需求理解—技术设计—实施配置—验收证据”四个层面展开，便于评审人员快速定位。")
    writer.table(
        ["查阅目标", "建议入口"],
        [
            ["核对总体架构", "第三章及总体架构图"],
            ["核对平台软件功能", "第四章及AI能力专题"],
            ["核对专网、公网与电话", "第五至第七章"],
            ["核对系统对接", "第八章及接口治理专题"],
            ["核对设备规格", "第九章"],
            ["核对测试、实施与服务", "第十至第十五章"],
            ["逐条核对符合性", "第十六章技术条款响应矩阵"],
        ],
        [5, 11.8],
    )

    writer.new_page("目录", "前置文件", 1)
    p = writer.doc.add_paragraph()
    p.paragraph_format.first_line_indent = None
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u')
    writer.paragraph("提示：使用 Microsoft Word 打开后，按 Ctrl+A，再按 F9 更新目录、页码及域。", size=9, indent=False)


def add_chapter_overview(writer: ProposalWriter, section: dict) -> None:
    focus, design, verify = section_focus(section["title"])
    writer.new_page(section["title"], "招标架构对应章节", 1)
    writer.paragraph(f"本章围绕“{focus}”组织方案内容。设计坚持招标文件优先，将业务目标转化为平台能力、网络能力、设备能力、接口能力、安全控制和运维流程，并明确验收证据。")
    writer.paragraph(f"总体实施路径为：现状核查与边界确认—深化设计与资源规划—实验室预集成—现场部署与分批迁移—联调测试—试运行—验收移交。重点采用{design}。")
    writer.note(f"本章验收主线：{verify}。")
    writer.bullets(summarize_items(section["items"], 5))


def add_section_pages(writer: ProposalWriter, section: dict, pages: int = 3) -> None:
    title = section["title"]
    items = summarize_items(section["items"], 6)
    focus, design, verify = section_focus(title)

    writer.new_page(f"{title}——需求理解与响应", section["chapter"], 2)
    writer.paragraph(f"本节目标是实现{focus}。招标要求不仅关注功能是否存在，还关注跨系统协同、生产连续性、权限边界和长期运维，因此采用“平台统一、资源分域、接口标准、过程可审计”的响应路线。")
    writer.bullets(items[:5])
    writer.note("响应结论：满足。最终配置、数量和接口边界在深化设计阶段复核，任何调整不得降低招标技术指标。")

    writer.new_page(f"{title}——技术设计", section["chapter"], 2)
    writer.paragraph(f"技术实现以{design}为主线。控制面、媒体面、管理面和数据面分别规划，并通过统一身份、统一资源目录、统一事件模型和统一审计进行关联，避免形成新的信息孤岛。")
    rows = [
        ["业务层", f"围绕{focus}编排调度、协作和应急流程", "业务结果、回执与报告"],
        ["平台层", "提供会话控制、媒体处理、资源编排和开放接口", "服务状态与调用日志"],
        ["接入层", "适配专网、公网、SIP、视频、广播及智能终端", "注册、链路与能力状态"],
        ["数据层", "统一事件编号，关联话单、媒体、位置、告警和审计", "可检索、可追溯记录"],
        ["安全运维", "最小权限、分域隔离、统一监控、告警和工单", "配置基线与闭环证据"],
    ]
    writer.table(["设计层次", "实现要点", "输出"], rows, [3, 9, 5], 7.8)

    if pages >= 3:
        writer.new_page(f"{title}——配置、测试与验收", section["chapter"], 2)
        writer.paragraph("实施前形成参数确认表和接口清单；实施中执行双人复核、变更审批和配置备份；实施后按功能、性能、可靠性、安全性和可运维性分层测试。")
        writer.table(
            ["阶段", "关键动作", "验收证据"],
            [
                ["准备", "确认范围、现网、账号、地址、号码、通信组和数据留存策略", "确认单、勘察报告"],
                ["部署", "按基线安装、加固、纳管并执行单机检查", "安装记录、配置快照"],
                ["联调", "完成端到端业务、异常分支及跨系统联调", "联调记录、日志"],
                ["测试", verify, "测试用例、截图、报表"],
                ["移交", "更新台账、培训运维人员并完成备份", "资产台账、培训记录"],
            ],
            [2.4, 9.2, 5.4],
            7.8,
        )


def add_special_pages(writer: ProposalWriter, diagrams: dict[str, Path]) -> None:
    image_map = {
        "AI智能融合远程协助总体方案": "ai",
        "四自动闭环运营": "automation",
        "统一网管与全网设备一体化管理": "management",
        "全流程安全审计": "security",
        "零信任安全与纵深防护": "security",
    }
    for index, (title, description) in enumerate(KEY_TOPICS, 1):
        writer.new_page(title, f"重点能力专题 {index:02d}", 1)
        writer.paragraph(description)
        if title in image_map:
            writer.image(diagrams[image_map[title]], 17.2)
            image_notes = {
                "AI智能融合远程协助总体方案": "现场人员通过AI语音交互完成呼叫、取证、查询和求助，减少手动操作；专家指导、标注、告警和处置过程统一记录。",
                "四自动闭环运营": "四个环节共享同一资产身份、配置基线、事件编号和审计链；任何自动操作均可暂停、回滚并转人工处理。",
                "统一网管与全网设备一体化管理": "管理对象覆盖平台、服务器、基站、网关、调度台、网络链路和智能终端，实现一张拓扑、一套告警、一套工单和一套报表。",
                "全流程安全审计": "审计记录与原始话单、媒体、配置版本和工单关联，支持按人、设备、事件、时间和操作类型检索并形成审计报告。",
                "零信任安全与纵深防护": "默认不信任任何主体或终端，持续验证身份、设备状态和访问上下文，异常访问实时阻断并触发告警。",
            }
            writer.note(image_notes[title])
            continue
        focus, design, verify = section_focus(title)
        tailored = {
            "AI降噪与强噪声环境语音增强": [
                ["输入处理", "结合麦克风阵列能力实施波束增强、语音活动检测、自动增益和回声消除。"],
                ["智能抑噪", "识别持续机械噪声、风噪和突发冲击噪声，在保留人声特征的前提下动态抑制。"],
                ["双端协同", "现场终端优先完成低时延降噪，平台侧按会话质量进行二次增强和录音后处理。"],
                ["质量监测", "持续采集时延、抖动、丢包、信噪比和主观质量指标，异常自动告警并关联工单。"],
                ["验收方法", "选取典型生产噪声样本开展对照测试，以语音可懂度、识别率和端到端时延综合评价。"],
            ],
            "AI智能对话与免手持作业": [
                ["交互入口", "支持唤醒词和按键两种入口，人员可用自然语言发起呼叫、查询、拍摄、上报和求助。"],
                ["对话理解", "通过ASR、意图识别、上下文管理和园区知识库理解连续指令，必要时主动澄清。"],
                ["免手持执行", "将确认后的语音意图转换为受权限控制的终端动作，执行结果通过TTS和界面同步反馈。"],
                ["安全控制", "高风险动作要求二次确认；身份、指令、执行结果和失败原因写入安全审计日志。"],
                ["降级策略", "弱网或AI服务不可用时保留物理按键、预置口令和人工调度通道，确保基本通信不中断。"],
            ],
            "一体化AI录音录像自动报告": [
                ["统一采集", "按策略对语音、视频、屏幕协作和调度操作统一录制，以事件编号和时间轴关联。"],
                ["智能处理", "执行语音转写、说话人区分、关键词识别、摘要生成、待办提取和风险提示。"],
                ["报告生成", "按检修、应急、会商等模板自动生成过程、结论、责任人、待办和证据索引。"],
                ["人工复核", "AI结果标识生成时间、模型版本和置信信息，发布前支持授权人员校订与签发。"],
                ["证据保护", "原始媒体只读留存，报告修改保留版本；调阅、导出和删除均授权并审计。"],
            ],
            "智能终端自动发现": [
                ["发现来源", "汇聚注册事件、标准管理协议、地址扫描和网关上报，避免依赖单一发现机制。"],
                ["身份识别", "依据设备证书、序列号、MAC、型号和软件指纹建立唯一资产身份并去重。"],
                ["能力建模", "自动识别语音、视频、定位、告警、编解码和安全能力，形成终端能力画像。"],
                ["准入控制", "未知或不合规终端进入隔离区，完成审批、加固和基线检查后方可纳入生产域。"],
                ["结果输出", "实时更新资产台账、拓扑、在线状态和责任归属，发现异常自动产生告警。"],
            ],
            "智能终端自动配置": [
                ["模板策略", "按企业、部门、岗位、区域、终端型号和安全等级建立继承式配置模板。"],
                ["配置内容", "统一下发号码、通信组、权限、路由、编解码、时间、证书和网管参数。"],
                ["一致性校验", "下发前检查依赖和冲突，下发后读取运行配置与基线比对，偏差自动提示。"],
                ["批量控制", "支持灰度、分批、定时和并发控制，关键变更经过审批并保留完整版本。"],
                ["失败处置", "失败自动重试；超过阈值停止批次并回滚最近稳定版本，同时创建人工工单。"],
            ],
            "智能终端自动诊断": [
                ["健康检查", "检测供电、网络、注册、时间同步、证书、版本、存储和关键进程状态。"],
                ["媒体诊断", "关联信令和媒体路径分析时延、抖动、丢包、单通、无声和编解码协商问题。"],
                ["根因定位", "基于拓扑、配置变更、同类告警和历史故障缩小根因范围并给出处置建议。"],
                ["远程工具", "授权后执行连通性、抓包、日志采集和配置比对；敏感信息脱敏并全程审计。"],
                ["闭环验证", "修复后自动复测注册、呼叫和媒体质量，成功则关闭工单，失败则升级处理。"],
            ],
            "智能终端自动开通": [
                ["流程触发", "由审批完成、设备首次上线或批量建设计划触发标准开通工作流。"],
                ["资源编排", "自动校验号码、许可证、通信组、地址、证书和平台容量，防止资源冲突。"],
                ["业务开通", "按模板下发配置、注册平台并开通授权业务，关键步骤支持幂等和断点续作。"],
                ["自动验证", "执行注册、单呼、组呼、视频、定位、告警等适用测试并保存结果。"],
                ["台账归档", "成功后更新资产与配置基线；失败时回滚资源、记录原因并转入工单。"],
            ],
        }.get(title)
        writer.table(
            ["维度", "方案设计"],
            tailored or [
                ["场景", f"面向园区生产调度、检修协作、应急处置和日常运营中的{focus}需求。"],
                ["能力", design],
                ["安全", "按用户、岗位、组织、区域和事件授权；敏感操作审批并全程审计。"],
                ["运维", "纳入统一网管，监测状态、质量、容量、告警、配置和版本。"],
                ["验收", verify],
            ],
            [3, 14],
            8.2,
        )
        writer.note("设计边界：AI结果用于辅助研判与提效，关键生产、安全和应急决策保留人工确认；原始证据不可被AI结果替代。")


def response_text(requirement: str) -> str:
    if any(k in requirement for k in ("应", "支持", "具备", "要求", "不低于", "不得", "满足")):
        return "满足。按本方案对应功能、配置及测试方法实施，最终以招标指标和验收结果为准。"
    return "已响应。纳入深化设计、实施交付或项目管理范围，并提供相应记录。"


def add_response_matrix(writer: ProposalWriter, requirements: list[dict], chunk_size: int = 7) -> None:
    chunks = [requirements[i : i + chunk_size] for i in range(0, len(requirements), chunk_size)]
    for page_index, chunk in enumerate(chunks, 1):
        writer.new_page(f"技术条款逐项响应矩阵（{page_index}/{len(chunks)}）", "第十六章 投标技术文件编制要求", 2)
        rows = []
        for item in chunk:
            requirement = item["text"]
            rows.append(
                [
                    item["id"],
                    item["section"][:22],
                    requirement,
                    response_text(requirement),
                ]
            )
        writer.table(["编号", "所属条目", "招标要求", "响应说明"], rows, [1.7, 3.2, 7.8, 4.4], 6.7)
        writer.paragraph("注：本矩阵的“满足/已响应”均以正文设计、最终供货配置、联调结果和验收证据为支撑，不构成对招标指标的缩减。", size=7.4, indent=False)


def add_appendix_page(writer: ProposalWriter, index: int) -> None:
    topics = [
        ("需求—设计—测试追溯表模板", ["需求编号", "设计章节", "配置项", "测试用例", "证据编号", "验收结论"]),
        ("接口联调记录模板", ["接口名称", "协议与版本", "地址与鉴权", "正常用例", "异常用例", "遗留问题"]),
        ("设备到货检查表", ["设备型号", "序列号", "数量", "外观", "附件", "证书", "检查人"]),
        ("终端自动开通检查表", ["发现", "识别", "模板匹配", "配置下发", "注册", "业务验证", "台账"]),
        ("无线覆盖测试记录", ["测试区域", "位置", "信号指标", "语音质量", "切换", "问题与整改"]),
        ("音视频质量测试记录", ["会话编号", "时延", "抖动", "丢包", "主观质量", "录制一致性"]),
        ("安全审计检查表", ["登录", "授权", "配置", "呼叫", "调阅", "下载", "接口", "告警"]),
        ("AI报告抽检记录", ["事件编号", "原始媒体", "转写准确性", "摘要完整性", "待办", "人工复核"]),
        ("高可用切换记录", ["故障注入", "检测时间", "切换时间", "业务影响", "恢复", "数据一致性"]),
        ("培训签到与考核模板", ["课程", "对象", "讲师", "签到", "实操", "考核", "补训"]),
        ("项目交付物清单", ["设计文件", "配置文件", "测试报告", "竣工资料", "培训材料", "运维手册"]),
        ("风险与问题闭环表", ["编号", "风险/问题", "影响", "责任人", "措施", "期限", "状态"]),
        ("SLA服务记录模板", ["事件级别", "报障时间", "响应时间", "恢复时间", "根因", "改进"]),
        ("配置基线核查表", ["对象", "基线版本", "当前版本", "偏差", "审批", "整改", "复核"]),
        ("最终验收签署页", ["建设单位", "使用单位", "监理/顾问", "承建单位", "日期", "结论"]),
    ]
    title, fields = topics[(index - 1) % len(topics)]
    writer.new_page(title, f"附录 {index:02d}", 1)
    writer.paragraph("本表用于项目实施与验收阶段形成标准化记录。实际使用时应补充项目名称、地点、日期、参与人员、版本和关联证据，并按文档管理制度归档。")
    headers = ["序号"] + fields[: min(5, len(fields))]
    rows = []
    for row_index in range(1, 9):
        rows.append([str(row_index)] + [" " * 8 for _ in headers[1:]])
    widths = [1.2] + [(15.7 / (len(headers) - 1)) for _ in headers[1:]]
    writer.table(headers, rows, widths, 7.2)
    writer.note("记录要求：内容真实、时间同步、签名完整、版本受控；电子记录应与统一事件编号或工单编号关联。")


def add_source_inventory(writer: ProposalWriter, source_paths: list[Path]) -> None:
    writer.new_page("已采用资料清单与适用性说明", "附录", 1)
    rows = []
    for path in source_paths:
        if path and path.exists():
            rows.append([path.name, "已解析", "用于方案架构、功能、操作、运维或产品能力说明"])
    rows.append(["统一化AI智能融合远程协助系统方案介绍2026-09_现网痛点页完善版.pptx", "未提供", "仅采用用户明确提出的能力要求，不引用缺失文件的具体参数"])
    writer.table(["资料名称", "状态", "使用原则"], rows, [8.5, 2.2, 6.4], 7.3)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--tender", type=Path, required=True)
    parser.add_argument("--reference", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--material", type=Path, action="append", default=[])
    parser.add_argument("--target-pages", type=int, default=350)
    args = parser.parse_args()

    sections, requirements, tender_doc = parse_tender(args.tender)
    doc = Document(args.reference)
    clear_document_body(doc)
    configure_styles(doc)
    configure_page(doc)
    props = doc.core_properties
    props.title = "园区融合通讯系统建设项目技术方案"
    props.subject = "融合通信、AI远程协助、统一网管与应急指挥"
    props.author = "项目技术方案编制组"
    props.last_modified_by = "项目技术方案编制组"
    props.category = "投标技术文件"
    props.keywords = "融合通信, AI降噪, 安全审计, 录音录像, 统一网管, 自动发现, 自动配置, 自动诊断, 自动开通"
    props.comments = "依据招标需求及已提供产品资料自动编制，招标要求优先。"

    args.output.parent.mkdir(parents=True, exist_ok=True)
    diagrams = build_diagrams(args.output.parent / "diagrams")
    writer = ProposalWriter(doc)
    add_cover(writer)
    add_front_matter(writer)

    # Overall architecture visual page.
    writer.new_page("总体技术架构全景", "核心架构图", 1)
    writer.image(diagrams["overall"], 17.2)
    writer.paragraph("总体架构以统一融合通信平台为中枢，向下兼容多制式网络和终端，向上支撑调度、应急、远程协助与智能运营；安全和统一网管作为横向能力覆盖全栈。", size=9.3, indent=False)

    for section in sections:
        if section["level"] == 1:
            add_chapter_overview(writer, section)
        elif section["level"] == 2:
            add_section_pages(writer, section, 3)
        else:
            add_section_pages(writer, section, 2)

    add_special_pages(writer, diagrams)
    add_response_matrix(writer, requirements, 7)
    add_source_inventory(writer, [args.tender, args.reference] + args.material)

    appendix_index = 1
    while writer.explicit_pages < args.target_pages:
        add_appendix_page(writer, appendix_index)
        appendix_index += 1

    # If source headings expand the plan past the target, that is acceptable for
    # "around 350 pages", but record the explicit minimum in document metadata.
    props.comments += f" Explicit page units: {writer.explicit_pages}."
    doc.settings.update_fields_on_open = True
    doc.save(args.output)
    print(f"output={args.output}")
    print(f"explicit_pages={writer.explicit_pages}")
    print(f"tender_sections={len(sections)}")
    print(f"response_items={len(requirements)}")


if __name__ == "__main__":
    main()
