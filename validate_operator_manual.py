#!/usr/bin/env python3
"""Validate the generated operator manual."""

from __future__ import annotations

import argparse
import subprocess
import tempfile
import zipfile
from pathlib import Path

from docx import Document


REQUIRED_TEXT = [
    "AI智能远程协助系统",
    "语音调度",
    "视频会议",
    "会议会商",
    "录音录像",
    "状态服务已连接",
    "AI能力与自动化功能边界",
    "安全、审计与数据保护",
    "故障处理",
    "班次交接",
]

FORBIDDEN_TEXT = [
    "123456",
    "为了凑页",
    "待补充内容",
    "TODO",
]


def pdf_page_count(docx_path: Path) -> int:
    with tempfile.TemporaryDirectory(prefix="operator-manual-") as tmp:
        subprocess.run(
            [
                "libreoffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                tmp,
                str(docx_path),
            ],
            check=True,
            capture_output=True,
            text=True,
            timeout=180,
        )
        pdf = Path(tmp) / f"{docx_path.stem}.pdf"
        result = subprocess.run(
            ["pdfinfo", str(pdf)],
            check=True,
            capture_output=True,
            text=True,
            timeout=30,
        )
        for line in result.stdout.splitlines():
            if line.startswith("Pages:"):
                return int(line.split(":", 1)[1].strip())
    raise RuntimeError("Unable to determine PDF page count")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("document", type=Path)
    args = parser.parse_args()
    path = args.document.resolve()
    if not path.exists():
        raise SystemExit(f"Document not found: {path}")

    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member:
            raise SystemExit(f"Corrupt DOCX member: {bad_member}")

    doc = Document(path)
    paragraph_text = "\n".join(p.text for p in doc.paragraphs)
    table_text = "\n".join(
        cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    )
    all_text = f"{paragraph_text}\n{table_text}"

    missing = [text for text in REQUIRED_TEXT if text not in all_text]
    forbidden = [text for text in FORBIDDEN_TEXT if text in all_text]
    if missing:
        raise SystemExit(f"Missing required text: {missing}")
    if forbidden:
        raise SystemExit(f"Forbidden text present: {forbidden}")

    body_paragraphs = [p for p in doc.paragraphs if p.text.strip()]
    if len(body_paragraphs) < 180:
        raise SystemExit(f"Too few non-empty paragraphs: {len(body_paragraphs)}")
    if len(doc.tables) < 15:
        raise SystemExit(f"Too few tables: {len(doc.tables)}")
    if len(doc.inline_shapes) < 12:
        raise SystemExit(f"Too few screenshots: {len(doc.inline_shapes)}")

    pages = pdf_page_count(path)
    if not 25 <= pages <= 80:
        raise SystemExit(f"Unexpected rendered page count: {pages}")

    print(f"document={path}")
    print(f"pages={pages}")
    print(f"paragraphs={len(body_paragraphs)}")
    print(f"tables={len(doc.tables)}")
    print(f"images={len(doc.inline_shapes)}")
    print(f"text_chars={len(all_text)}")
    print("status=PASS")


if __name__ == "__main__":
    main()
