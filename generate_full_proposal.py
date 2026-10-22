#!/usr/bin/env python3
"""Build a continuous, narrative tender proposal from reviewed chapter drafts."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from generate_proposal import (
    BLUE,
    FONT_NAME,
    GRAY,
    LIGHT_BLUE,
    MID_BLUE,
    PALE_BLUE,
    ProposalWriter,
    add_field,
    build_diagrams,
    clean_text,
    clear_document_body,
    configure_page,
    configure_styles,
    create_diagram,
    parse_tender,
    set_cell_margins,
    set_cell_shading,
    set_repeat_table_header,
    set_run_font,
)

SOLUTION_NAME = "基于“云+管+端”架构一体化AI智能融合调度指挥系统"


SCENARIOS = [
    ("园区总调度日常生产指挥", "园区总调度面向多个企业和约80个部门组织日常生产协同，需要在不突破企业管理边界的前提下掌握全局通信资源。", "值班员在统一调度台查看组织、通信组和终端状态，按岗位或区域发起单呼、组呼、会议或广播；跨企业任务经授权后建立临时通信组，结束后自动解散并归档。"),
    ("炼钢生产异常跨部门协同", "设备、能源、生产、安全等专业同时参与处置，传统逐级电话通知容易遗漏且信息不一致。", "告警触发预案后，系统按责任矩阵并行通知相关岗位，自动统计接通和确认结果；未确认人员按策略重试或升级，指挥员可随时加入视频、监控画面和GIS位置。"),
    ("高噪声区域现场检修", "风机、泵组、车辆和金属碰撞产生持续及突发噪声，影响语音可懂度，也增加手动操作风险。", "智能终端启用AI降噪、回声消除和自动增益，作业人员通过唤醒词发起专家协助；远端专家查看第一视角视频并进行语音、标注和资料推送，作业全过程录音录像。"),
    ("智能安全帽免手持协作", "登高、持械或佩戴防护手套时无法频繁操作屏幕，需要语音完成常用任务。", "现场人员通过AI智能对话完成呼叫、拍照、录像、查询步骤、上报告警和请求支援；高风险动作二次确认，离线时降级到物理按键和预置指令。"),
    ("有限空间作业监护", "有限空间作业需要持续掌握人员位置、生命体征、环境告警和通信状态。", "作业审批形成临时任务组，人员进入区域后位置和终端状态持续上报；异常触发声光、语音和调度告警，监护人与应急人员一键入会，事件证据自动归档。"),
    ("高处作业跌落与脱帽告警", "高处作业的跌落、长时间静止或脱帽需要快速识别，误报也必须可核验。", "终端产生告警后平台关联人员、岗位、位置和附近视频，先执行语音确认；无应答或确认危险时自动升级，调度就近力量并记录完整处置时间线。"),
    ("港口装卸与车辆调度", "港口区域跨度大、移动性强，专网和公网覆盖条件不同，车辆与人员需要统一协同。", "系统根据终端能力和网络状态选择可用通信路径，GIS展示车辆和作业组位置；调度员可框选区域发起组呼，弱网时优先保障语音和关键消息。"),
    ("仓储作业与紧急广播", "仓储区域需要分区通知、人员疏散和广播回执，且应避免无关区域受到干扰。", "调度员选择仓库、楼层或自定义区域触发广播，系统联动电话、对讲、APP和广播终端；通知内容、播放次数、失败清单及人工补呼结果统一留痕。"),
    ("跨制式专公网对讲", "专网数字集群、公网对讲和IP电话彼此独立会形成协同断点。", "融合平台建立跨网通信组和资源映射，完成单呼、组呼、强插、强拆和会议桥接；边界控制限制跨域可见范围，并对跨网操作单独审计。"),
    ("既有IP电话平滑迁移", "生产电话不能因平台建设大规模停机，号码和使用习惯应尽量保留。", "先完成现网号码、路由、中继和终端摸排，在实验环境验证SIP兼容；按部门分批迁移并设置回退窗口，迁移后执行内外线、转接、录音和应急呼叫测试。"),
    ("视频监控联动调度", "事件处置时人工在多个平台查找摄像机耗时，且画面与通信事件缺少关联。", "平台通过标准接口获取摄像机目录，按GIS位置、告警点或终端位置推荐画面；授权用户可预览、轮巡、上墙和控制云台，调用行为进入审计。"),
    ("应急会商与多媒体会议", "重大事件需要园区、企业、现场和外部专家快速加入同一会商。", "预案一键创建会议，电话、对讲、视频终端、APP和会议室按能力接入；主席控制发言、静音、双流和录制，会议结束自动形成参会、媒体和待办记录。"),
    ("断网场景本地自治", "园区骨干或上联中断时，局部生产区域仍需保持基本通信。", "边缘节点保留必要注册、短号、组呼和应急号码能力，终端按预案切换可用链路；网络恢复后同步事件和审计记录，冲突数据按时间及权威源处理。"),
    ("断电与设备故障保障", "机房供电、单台服务器或链路故障不能导致核心通信长时间中断。", "核心服务冗余部署并接入后备电源，网管持续执行健康检查；故障触发切换、告警和工单，恢复后核对会话、配置与数据一致性。"),
    ("重大事件批量通知", "大规模通知需要快速触达并掌握人员确认情况，单纯广播无法确认个体状态。", "系统按组织、岗位、区域或名单并行发起电话、短信、APP和对讲通知，记录接通、已读和确认；失败对象按策略重试并生成未触达清单。"),
    ("敏感通话与反诈审计", "关键业务通话需要留证，异常号码、敏感词和可疑行为需要及时发现。", "按授权策略录音并执行智能转写和规则分析，命中风险时通知审计人员；原始录音、分析结果、处置意见和导出操作全程可追溯。"),
    ("AI录音录像自动报告", "远程协助和应急会商结束后，人工整理过程记录耗时且容易遗漏。", "平台关联原始媒体、转写、说话人、标注和操作轨迹，提取摘要、关键事件、结论和待办，生成结构化报告草稿；授权人员复核签发，AI结果不替代原始证据。"),
    ("智能终端批量开通", "终端数量增长后，逐台配置易出错且难以保持安全基线一致。", "终端上线后自动发现并识别能力，网管按组织和型号匹配模板，完成号码、通信组、权限和证书下发；自动验证注册与业务，失败则回滚并创建工单。"),
    ("智能终端故障自诊断", "一线报障通常只描述“不能打电话”，需要快速区分终端、网络、平台和配置问题。", "网管关联注册、链路、信令、媒体质量、配置变更和同类告警，自动执行检测并给出根因候选；修复后复测，结果与工单一并归档。"),
    ("多语言值守与本地适配", "园区存在中文、英文和印尼文使用需求，界面和告警表达必须一致。", "平台和终端按角色选择语言，关键术语、预案和提示语经统一词库管理；上线前由业务和当地人员复核翻译、法规及认证适用性。"),
]


def build_extra_diagrams(output_dir: Path) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    specs = {
        "deployment": (
            "园区多级部署与接入关系",
            [
                ("中心核心区", ["融合通信平台", "AI能力中心", "统一网管", "录音录像存储"]),
                ("企业分调度", ["分级调度台", "本地自治", "权限与通信组", "应急预案"]),
                ("网络接入", ["专网数字集群", "公网4G/5G", "园区IP网络", "SBC/安全边界"]),
                ("业务终端", ["对讲终端", "IP电话", "智能安全帽", "APP/视频终端"]),
            ],
            "中心统一资源和策略，企业按域自治；边缘保留必要通信能力，链路恢复后同步记录。",
        ),
        "integration": (
            "存量系统利旧与标准接口架构",
            [
                ("既有资源", ["IP电话", "视频监控", "应急广播", "GIS/业务系统"]),
                ("接入适配", ["SIP/SBC", "GB/T 28181", "广播网关", "REST API/SDK"]),
                ("融合平台", ["资源目录", "会话与媒体", "事件总线", "权限与审计"]),
                ("融合应用", ["统一调度", "应急指挥", "远程协助", "运维与报表"]),
            ],
            "通过标准协议、资源映射和接口治理实现利旧，接口异常不扩散到核心业务。",
        ),
        "ha": (
            "核心平台高可用与数据保护",
            [
                ("接入冗余", ["双链路", "边界双归属", "健康检查", "流量切换"]),
                ("服务冗余", ["无单点部署", "状态同步", "故障隔离", "分级降级"]),
                ("数据保护", ["事务一致性", "备份校验", "异地副本", "恢复演练"]),
                ("运营保障", ["统一监控", "告警关联", "自动工单", "复盘改进"]),
            ],
            "高可用目标通过架构、配置、监测和演练共同实现，切换后核验业务与数据一致性。",
        ),
        "ai-data": (
            "AI录音录像报告数据流程",
            [
                ("证据采集", ["语音/视频", "屏幕与标注", "位置与告警", "调度操作"]),
                ("智能处理", ["降噪增强", "ASR转写", "要点识别", "摘要与待办"]),
                ("受控存储", ["原始媒体", "结构化索引", "报告版本", "留存与备份"]),
                ("授权使用", ["检索回放", "人工复核", "审计取证", "知识沉淀"]),
            ],
            "AI结果与原始证据分层保存；模型版本、处理时间、人工修改和导出动作全程留痕。",
        ),
        "migration": (
            "现网平滑迁移与回退路径",
            [
                ("准备阶段", ["资源普查", "兼容验证", "号码与路由", "回退基线"]),
                ("试点阶段", ["小范围接入", "双轨验证", "问题整改", "用户确认"]),
                ("分批迁移", ["窗口审批", "配置备份", "业务割接", "即时验证"]),
                ("稳定移交", ["试运行", "性能观察", "资料更新", "正式验收"]),
            ],
            "每批迁移均设置停止条件和回退步骤，生产业务验证通过后才进入下一批次。",
        ),
    }
    result = {}
    for key, (title, columns, footer) in specs.items():
        path = output_dir / f"{key}.png"
        create_diagram(path, title, columns, footer)
        result[key] = path
    return result


def set_body_style(doc: Document, font_size: float = 11.0, line_spacing: float = 1.55) -> None:
    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    style.font.size = Pt(font_size)
    style.paragraph_format.line_spacing = line_spacing
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.first_line_indent = Cm(0.74)
    for section in doc.sections:
        section.top_margin = Cm(2.3)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.3)

    if "Heading 4" not in [s.name for s in doc.styles]:
        h4 = doc.styles.add_style("Heading 4", WD_STYLE_TYPE.PARAGRAPH)
    else:
        h4 = doc.styles["Heading 4"]
    h4.font.name = FONT_NAME
    h4._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    h4.font.size = Pt(11.5)
    h4.font.bold = True
    h4.font.color.rgb = RGBColor.from_string(MID_BLUE)
    h4.paragraph_format.space_before = Pt(7)
    h4.paragraph_format.space_after = Pt(4)
    h4.paragraph_format.keep_with_next = True


def set_document_headers(doc: Document) -> None:
    for section in doc.sections:
        header = section.header
        for paragraph in header.paragraphs:
            paragraph.clear()
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(f"{SOLUTION_NAME}｜技术方案")
        set_run_font(r, 7.5, False, GRAY)


def add_cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("基于“云+管+端”架构")
    set_run_font(r, 25, True, BLUE)
    p.paragraph_format.space_after = Pt(8)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("一体化AI智能融合调度指挥系统")
    set_run_font(r, 30, True, MID_BLUE)
    p.paragraph_format.space_after = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("投标技术方案")
    set_run_font(r, 22, True, BLUE)
    p.paragraph_format.space_after = Pt(26)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("融合通信｜智能调度｜AI远程协助｜统一运维")
    set_run_font(r, 15, False, GRAY)
    for _ in range(10):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("技术文件")
    set_run_font(r, 13, True, BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("二〇二六年十月")
    set_run_font(r, 12, False, GRAY)
    doc.add_page_break()


def add_front_matter(doc: Document) -> None:
    writer = ProposalWriter(doc)
    writer.started = True
    writer.explicit_pages = 1
    p = doc.add_paragraph("文档说明", style="Heading 1")
    p.paragraph_format.page_break_before = False
    writer.paragraph(
        "本方案依据《园区融合通讯系统技术招标需求书》编制，章节架构与招标文件保持对应。方案以满足招标全部技术要求为基本原则，在此基础上结合已提供的US2000产品白皮书、管理员手册、AI与IMS融合资料、远程协助资料、录音方案和UC2000资料，对系统架构、功能、接口、实施、验收和运维进行深化设计。"
    )
    writer.paragraph(
        f"整体方案命名为《{SOLUTION_NAME}》。“云”侧集中承载融合通信、AI能力、数据与业务应用；“管”侧统一组织公网、专网、固网、无线、卫星及安全运维体系；“端”侧覆盖调度台、IP电话、对讲终端、智能安全帽、移动APP及其他智能终端。三者通过统一身份、资源目录、事件模型和安全审计形成一体化闭环。"
    )
    writer.paragraph(
        "招标需求书是本方案的最高技术依据。产品资料中存在历史版本、行业示例或与本项目无关的场景时，仅提取已经明确且与本项目适用的能力，不直接沿用旧参数。任何产品参数与招标要求不一致时，以招标要求和最终经确认的深化设计为准。"
    )
    writer.paragraph(
        "用户指定的《统一化AI智能融合远程协助系统方案介绍2026-09_现网痛点页完善版.pptx》未实际提供，因此本方案不引用该文件中无法核验的页面或参数；AI降噪、AI智能对话、免手持操作、安全审计、一体化AI录音录像自动报告、统一网管和智能终端“四自动”等明确要求，已依据现有资料和招标边界进行完整设计。"
    )
    writer.note("阅读提示：正文采用连续论述方式；技术参数、设备清单和逐条响应矩阵集中放在相应章节及附录，不以表格替代方案设计。")
    doc.add_page_break()

    p = doc.add_paragraph("总体响应声明", style="Heading 1")
    p.paragraph_format.page_break_before = False
    for lead, text in [
        ("完整响应。", "本方案覆盖招标文件的总体架构、平台软件、专网数字集群、公网对讲、IP电话、视频、会议、广播、GIS、录音录像、应急指挥、智能APP、分级调度、运行管理、统一网管、开放接口、主要设备、测试验收、施工、培训和售后服务要求。"),
        ("开放兼容。", "系统优先采用SIP、GB/T 28181、REST API等标准协议和开放接口，通过边界会话控制、协议适配和接口治理实现现有资源利旧，避免形成新的封闭孤岛。"),
        ("安全可靠。", "核心服务按照冗余、高可用和可恢复原则设计，安全控制覆盖身份、终端、网络、应用、接口、数据和运维全过程，并通过安全审计形成可追溯证据链。"),
        ("面向运营。", "建设范围不仅包括平台上线，还包括资产、拓扑、性能、告警、配置、版本、日志、工单和知识库的统一管理，以及智能终端自动发现、自动配置、自动诊断、自动开通的闭环运营。"),
        ("可测可验。", "每项功能均在正文中说明实现机制、异常处理和验收思路，逐条响应矩阵将招标要求与方案章节、实施活动和测试证据关联。"),
    ]:
        writer.paragraph(lead + text, bold_lead=lead)
    doc.add_page_break()

    doc.add_paragraph("目录", style="Heading 1")
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = None
    add_field(p, 'TOC \\o "1-4" \\h \\z \\u')
    p = doc.add_paragraph("使用 Microsoft Word 打开文档后，按 Ctrl+A、F9 更新目录及页码。")
    p.paragraph_format.first_line_indent = None
    for run in p.runs:
        set_run_font(run, 9, False, GRAY)
    doc.add_page_break()


def add_paragraph(doc: Document, text: str, bullet: bool = False, numbered: bool = False) -> None:
    text = re.sub(r"(\*\*|__|`)", "", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    p = doc.add_paragraph()
    if bullet or numbered:
        p.paragraph_format.left_indent = Cm(0.72)
        p.paragraph_format.first_line_indent = Cm(-0.42)
        prefix = "• " if bullet else ""
        r = p.add_run(prefix + text)
    else:
        p.paragraph_format.first_line_indent = Cm(0.74)
        r = p.add_run(text)
    set_run_font(r, 11)


def flush_paragraph(doc: Document, buffer: list[str]) -> None:
    text = clean_text(" ".join(buffer))
    if text:
        add_paragraph(doc, text)
    buffer.clear()


def add_markdown(doc: Document, path: Path, first_chapter: bool = False) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    buffer: list[str] = []
    seen_h1 = False
    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            flush_paragraph(doc, buffer)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            flush_paragraph(doc, buffer)
            level = len(heading.group(1))
            title = clean_text(heading.group(2)).strip("*")
            if level == 1:
                if seen_h1 or not first_chapter:
                    doc.add_page_break()
                seen_h1 = True
            p = doc.add_paragraph(title, style=f"Heading {level}")
            if level == 1:
                p.paragraph_format.page_break_before = False
            continue
        bullet = re.match(r"^\s*[-*]\s+(.+)$", line)
        numbered = re.match(r"^\s*(\d+[.)、])\s*(.+)$", line)
        if bullet or numbered:
            flush_paragraph(doc, buffer)
            text = clean_text(bullet.group(1) if bullet else f"{numbered.group(1)} {numbered.group(2)}")
            add_paragraph(doc, text, bullet=bool(bullet), numbered=bool(numbered))
            continue
        if line.strip().startswith("|"):
            # Draft agents were instructed to avoid tables; retain any accidental
            # table row as readable prose rather than creating table-heavy pages.
            flush_paragraph(doc, buffer)
            cells = [clean_text(x) for x in line.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", x or "-") for x in cells):
                add_paragraph(doc, "；".join(x for x in cells if x))
            continue
        buffer.append(line)
    flush_paragraph(doc, buffer)


def add_figure_page(doc: Document, title: str, image: Path, caption: str) -> None:
    doc.add_page_break()
    doc.add_paragraph(title, style="Heading 1")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.add_run().add_picture(str(image), width=Cm(16.5))
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    for run in p.runs:
        set_run_font(run, 9, False, GRAY)


def add_user_scheme_figure(doc: Document, image: Path) -> None:
    doc.add_page_break()
    doc.add_paragraph("多业务融合指挥调度方案架构", style="Heading 1")
    add_paragraph(
        doc,
        "本架构图从业务呈现和端到端连接角度描述园区融合通讯系统。系统通过统一支撑层屏蔽不同网络、协议和终端的差异，向上为调度台、业务融合平台、桌面终端和移动单兵APP提供一致的通信、调度、数据与管理能力。",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.add_run().add_picture(str(image), width=Cm(16.8))
    p = doc.add_paragraph("图：多业务融合指挥调度端到端架构（用户提供）")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    for run in p.runs:
        set_run_font(run, 9, False, GRAY)

    doc.add_paragraph("架构分层说明", style="Heading 2")
    add_paragraph(
        doc,
        "接入层承载监控、IP话机、视频会议、音响、麦克风、手机、车载台、对讲机、卫星终端、一键通和传真机等多类终端。终端是否纳入本次供货以招标清单为准，但平台在协议和资源模型上预留统一接入、状态管理和权限控制能力。",
    )
    add_paragraph(
        doc,
        "传输层兼容移动公网、移动专网、卫星通信、园区固网公网和固网专网等网络。融合平台不假定单一网络始终可用，而是结合业务等级、终端能力和网络状态选择可用路径；网络异常时按照语音优先、关键指令优先和本地自治原则降级。",
    )
    add_paragraph(
        doc,
        "支撑层由统一通信平台、视频融合汇聚、综合语音接入、辅助调度、数据整合分析、GIS地理整合及标准开放接口组成。该层统一管理用户、号码、通信组、媒体资源、位置、告警和事件，并通过接口治理向既有及未来业务系统开放能力。",
    )
    add_paragraph(
        doc,
        "应用层提供统一通信录、语音调度、会议会商、视频调度、GIS调度、APP调度、短信、传真、录音录像和移动APP等业务。各业务共享同一身份、组织、权限、事件编号和审计体系，避免同一用户在多个孤立系统重复配置。",
    )
    add_paragraph(
        doc,
        "展现层面向统一调度台、业务应用系统平台、桌面调度终端和移动单兵APP提供人机交互。不同终端根据岗位和场景呈现适配界面，操作结果回写统一事件中心，并纳入统一网管和安全审计。",
    )


def add_function_expansion_figure(doc: Document, image: Path) -> None:
    doc.add_page_break()
    doc.add_paragraph("一体化融合协同功能体系", style="Heading 1")
    add_paragraph(
        doc,
        "基于“云+管+端”架构，系统通过受控网关和标准接口汇聚现网窄带、宽带、固网、移动公网、园区专网、有线与无线资源，在统一身份、组织、权限和事件模型下形成语音、视频、数据和终端管理四类融合协同能力。融合的目标不是把各专业系统简单集中展示，而是让调度动作能够跨系统执行、状态能够逐对象反馈、异常能够隔离降级、全过程能够审计追溯。",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.add_run().add_picture(str(image), width=Cm(16.8))
    p = doc.add_paragraph("图：一体化融合协同功能体系（用户提供）")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    for run in p.runs:
        set_run_font(run, 9, False, GRAY)

    doc.add_paragraph("语音融合协同", style="Heading 2")
    add_paragraph(
        doc,
        "语音融合协同以软交换、边界会话控制、数字集群接口和媒体网关为基础，将既有IP电话、PSTN或适用模拟线路、专网数字集群、公网POC、车载及手持对讲、移动APP、卫星语音和会议资源纳入统一号码及通信组体系。园区总调度与企业分调度共享同一资源目录，各企业保持用户、通信组、调度权限和数据范围独立；跨企业协同通过有期限、有范围的临时授权或应急预案实现。",
    )
    add_paragraph(
        doc,
        "调度员可以从组织树、通讯录、GIS地图或预案选择人员和终端，发起单呼、组呼、强插、强拆、监听、电话会议、跨网会商、语音通知或广播。平台负责SIP电话全双工会话与对讲半双工发言权之间的转换，记录申请、准许、占用、释放和超时状态，避免电话侧背景噪声或未挂机持续占用专网信道。专网、公网或电话接口中任一侧异常时，故障被限制在对应融合链路，专业系统内部基础通信继续运行。",
    )
    add_paragraph(
        doc,
        "语音质量设计覆盖编解码协商、回声抑制、抖动缓冲、丢包补偿、QoS和AI降噪。验收不仅检查界面是否出现按钮，还要使用真实终端完成双向语音、组呼、紧急呼叫、跨网发言控制、失败成员反馈、录音和话单关联，并在弱网、接口中断和节点切换条件下验证降级与恢复。",
    )

    doc.add_paragraph("视频融合协同", style="Heading 2")
    add_paragraph(
        doc,
        "视频融合协同在不替代既有视频监控平台管理职责的前提下，通过GB/T 28181及对端实际开放的标准接口同步监控目录、在线状态和能力标签，并接入视频会议终端、移动APP、智能安全帽、单兵终端及远程专家工作台。调度人员可按组织、区域、事件或GIS位置检索资源，执行实时预览、轮巡、上墙、云台控制、录像索引和视频入会。",
    )
    add_paragraph(
        doc,
        "应急处置时，系统把现场第一视角视频、附近监控画面、会议会商、语音通话、位置和告警关联到同一事件编号。远程专家可以冻结关键帧、标注、推送资料和形成建议卡片，现场人员逐项确认执行状态。弱网时优先保障语音和控制消息，视频自适应降低码率或退化为关键图片；视频故障不得影响专网语音、固定电话和应急广播。",
    )
    add_paragraph(
        doc,
        "视频目录可见、实时调阅、云台控制、录像查询、跨网分发和导出分别授权。所有调阅和控制动作记录用户、资源、时间、用途与结果。验收覆盖目录全量及增量同步、权限隔离、主辅流、多画面、移动视频、网络自适应、接口异常和恢复，使用真实画面及日志证明端到端效果。",
    )

    doc.add_paragraph("数据融合协同", style="Heading 2")
    add_paragraph(
        doc,
        "数据融合协同以统一资源目录和事件中心为枢纽，关联人员、组织、岗位、终端、号码、通信组、位置、视频、广播分区、告警、话单、录音录像、任务和工单。GIS、即时消息、OA、应急管理及其他业务系统通过受控API、SDK或事件订阅接入，数据主责仍由来源系统承担，融合平台保存稳定标识映射和必要业务副本，不通过复制全部数据库实现表面统一。",
    )
    add_paragraph(
        doc,
        "一次生产异常或应急事件从接报开始即生成唯一事件编号，后续预案匹配、人员通知、临时建组、语音视频会商、GIS调度、广播、远程协助、录音录像、任务反馈和AI报告全部关联该编号。接口采用鉴权、最小授权、限流、幂等、防重放、超时、重试和状态查询机制；请求超时时先查询真实执行状态，避免重复建会、重复广播或重复下发任务。",
    )
    add_paragraph(
        doc,
        "融合分析基于真实业务记录形成接通率、响应时长、任务闭环、资源使用、故障影响和趋势报表。数据缺失必须标记缺口，不能按零值处理；AI生成的摘要、待办和报告保存模型版本、输入来源和人工复核记录，原始媒体及日志保持不变。验收通过同一事件在来源系统、融合平台、审计记录和报告中的编号、时间、状态与结果交叉比对。",
    )

    doc.add_paragraph("终端统一监控、维护与管理", style="Heading 2")
    add_paragraph(
        doc,
        "统一网管对软交换、服务器、数据库、媒体服务、边界设备、基站、载波、网关、调度台、IP电话、专公网对讲终端、智能安全帽和移动APP建立统一资产及拓扑视图。每个对象保存厂家、型号、序列号、版本、地址、组织归属、位置、能力、配置基线和责任人；跨厂家管理遵循实际开放协议，无法自动执行的项目明确标识人工步骤，不虚构自动化覆盖。",
    )
    add_paragraph(
        doc,
        "自动发现通过注册事件、标准管理协议、地址扫描、网关上报和设备指纹识别新增终端，未知或不合规设备先进入隔离状态。自动配置根据企业、部门、岗位、区域、终端类型和安全等级匹配模板，下发号码、通信组、权限、编解码、时间、证书及网管参数，并回读运行配置与批准基线比对。",
    )
    add_paragraph(
        doc,
        "自动诊断关联注册、链路、信令、媒体质量、版本、证书、配置变化和同类告警，给出根因候选及处置建议；修复后自动复测，不以告警消失代替业务恢复。自动开通把审批、资源校验、号码分配、模板下发、注册、业务测试和台账更新编排为可暂停、可重试、可回滚的工作流，批量操作先试点、再灰度扩展，超过失败阈值自动熔断并转人工工单。",
    )
    add_paragraph(
        doc,
        "终端统一管理的验收以闭环结果为准：发现结果要进入资产台账，配置结果要完成回读，诊断结果要有证据和复测，开通结果要通过真实呼叫及适用业务验证。任何远程写配置、升级、遥晕遥毙、证书下发和批量操作均执行审批、范围校验和审计，失败时不得遗留冲突号码、孤立账号或半配置设备。",
    )


def add_scenarios(doc: Document) -> None:
    doc.add_page_break()
    doc.add_paragraph("典型业务场景深化设计", style="Heading 1")
    add_paragraph(
        doc,
        "本章通过典型生产和应急场景说明各子系统如何协同工作。场景设计不替代招标条款，而是把平台功能、网络能力、终端能力、安全控制和运维机制串联为端到端流程，便于深化设计、用户培训和验收演练。",
    )
    for index, (name, context, process) in enumerate(SCENARIOS, 1):
        doc.add_paragraph(f"{index}. {name}", style="Heading 2")
        add_paragraph(doc, f"场景背景与目标：{context}")
        add_paragraph(doc, f"业务流程：{process}")
        add_paragraph(
            doc,
            "控制与异常处理：业务动作依据组织、岗位和事件授权执行，跨域访问需显式授权；通信路径不可用时按照语音优先、关键消息优先和本地自治原则降级。自动化步骤失败后停止后续高风险动作，保留已完成步骤和失败原因，并转入人工工单。",
        )
        add_paragraph(
            doc,
            "记录与验收：系统保留会话、位置、告警、调度操作、配置变化和处置结果，涉及录音录像时按留存和调阅策略管理。验收采用正常流程、权限边界、网络异常和设备故障四类用例，核对业务结果、时间线、审计记录和报告的一致性。",
        )


def add_tender_tables(doc: Document, tender_doc: Document) -> None:
    doc.add_page_break()
    doc.add_paragraph("招标原始清单与数据表", style="Heading 1")
    add_paragraph(
        doc,
        "本附录原样整理招标需求书中的结构化清单，便于供货、深化设计和验收阶段核对。重复表头或合并单元格内容以原文为准；最终设备型号、数量和配置应通过投标澄清及合同文件确认。",
    )
    writer = ProposalWriter(doc)
    writer.started = True
    for index, source in enumerate(tender_doc.tables, 1):
        doc.add_paragraph(f"招标数据表 {index}", style="Heading 2")
        if not source.rows:
            continue
        headers = [clean_text(c.text) or f"字段{i+1}" for i, c in enumerate(source.rows[0].cells)]
        rows = [[clean_text(c.text) for c in row.cells] for row in source.rows[1:]]
        if rows:
            writer.table(headers, rows, size=7.5)


def add_response_matrix(doc: Document, requirements: list[dict]) -> None:
    doc.add_page_break()
    doc.add_paragraph("招标技术条款逐项响应矩阵", style="Heading 1")
    add_paragraph(
        doc,
        "响应矩阵用于证明正文与招标条款的完整对应关系，不替代前述详细方案。每项“满足”均以方案设计、最终配置、实施记录和测试证据为支撑；如合同澄清形成更高要求，按更高要求执行。",
    )
    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    for index, width in enumerate((1.5, 3.2, 7.2, 4.5)):
        table.columns[index].width = Cm(width)
    header = table.rows[0]
    set_repeat_table_header(header)
    headers = ["编号", "所属条目", "招标要求", "响应及落实方式"]
    for i, text in enumerate(headers):
        set_cell_shading(header.cells[i], BLUE)
        set_cell_margins(header.cells[i])
        p = header.cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, 7.2, True, "FFFFFF")
    for index, item in enumerate(requirements):
        cells = table.add_row().cells
        values = [
            item["id"],
            item["section"],
            item["text"],
            "满足。已纳入对应章节的设计、实施和验收范围；以招标指标为最低要求，并通过配置记录、测试结果及交付资料形成证据。",
        ]
        for i, value in enumerate(values):
            if index % 2:
                set_cell_shading(cells[i], PALE_BLUE)
            set_cell_margins(cells[i], 35, 45, 35, 45)
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(value)
            set_run_font(r, 6.5)

    doc.add_paragraph("技术响应结论", style="Heading 2")
    add_paragraph(
        doc,
        "本方案已将招标要求落实到总体架构、平台软件、专网数字集群、公网对讲、IP电话利旧、视频与广播融合、GIS调度、主要设备、实施、测试、培训和服务等完整建设链条。逐条响应矩阵用于建立要求与方案之间的索引关系，真正的技术响应由前述连续正文、最终配置、实施记录和测试证据共同构成，不以表格中的“满足”二字代替工程实现。",
    )
    add_paragraph(
        doc,
        "进入深化设计后，项目团队将以本矩阵为需求基线，为每一项条款补充设计对象、配置项、责任单位、测试用例和证据编号。现场条件、第三方接口或当地合规要求引起设计调整时，调整结果必须保持或提高招标指标，并通过正式评审和变更流程更新方案、清单、测试及交付资料，确保投标承诺能够连续传递到竣工验收。",
    )
    add_paragraph(
        doc,
        "对于AI降噪、智能对话、远程协助、录音录像自动报告、全流程安全审计、统一网管和终端“四自动”等增强能力，方案始终以生产通信可靠性为前提。智能能力可旁路、结果可复核、数据可追溯、故障可降级，任何模型或自动化流程都不会取代专网核心语音、现场安全制度和授权人员的最终判断。",
    )
    add_paragraph(
        doc,
        "通过统一的需求、设计、配置、测试和证据管理，本项目最终交付的不只是能够演示的功能集合，而是一套在园区复杂工业环境中可长期运行、可扩展、可维护、可审计的一体化AI智能融合调度指挥系统。",
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--tender", type=Path, required=True)
    parser.add_argument("--reference", type=Path, required=True)
    parser.add_argument("--draft", type=Path, action="append", required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--scheme-image", type=Path)
    parser.add_argument("--function-image", type=Path)
    parser.add_argument("--font-size", type=float, default=11.0)
    parser.add_argument("--line-spacing", type=float, default=1.55)
    args = parser.parse_args()

    sections, requirements, tender_doc = parse_tender(args.tender)
    missing = [str(path) for path in args.draft if not path.exists()]
    if missing:
        raise SystemExit("Missing draft files: " + ", ".join(missing))

    doc = Document(args.reference)
    clear_document_body(doc)
    configure_styles(doc)
    configure_page(doc)
    set_body_style(doc, args.font_size, args.line_spacing)
    set_document_headers(doc)
    props = doc.core_properties
    props.title = f"{SOLUTION_NAME}投标技术方案"
    props.subject = "完整融合通信技术方案"
    props.author = "项目技术方案编制组"
    props.last_modified_by = "项目技术方案编制组"
    props.category = "投标技术文件"
    props.keywords = "融合通信, 数字集群, 公网对讲, IP电话, AI远程协助, 安全审计, 统一网管, 四自动"
    props.comments = "正文型完整方案；招标需求优先。"
    doc.settings.update_fields_on_open = True

    args.output.parent.mkdir(parents=True, exist_ok=True)
    diagrams = build_diagrams(args.output.parent / "full-diagrams")
    extra_diagrams = build_extra_diagrams(args.output.parent / "full-diagrams")
    add_cover(doc)
    add_front_matter(doc)
    add_figure_page(doc, "园区融合通讯系统总体架构", diagrams["overall"], "图：统一接入、融合能力、业务应用和安全运维四层架构")
    if args.scheme_image:
        if not args.scheme_image.exists():
            raise SystemExit(f"Missing scheme image: {args.scheme_image}")
        add_user_scheme_figure(doc, args.scheme_image)
    if args.function_image:
        if not args.function_image.exists():
            raise SystemExit(f"Missing function image: {args.function_image}")
        add_function_expansion_figure(doc, args.function_image)

    for index, draft in enumerate(args.draft):
        if draft.name == "ch05-08.md":
            add_figure_page(doc, "园区多级部署与接入设计", extra_diagrams["deployment"], "图：中心、企业分调度、接入网络和业务终端的分层部署")
        if draft.name == "ch09-12.md":
            add_figure_page(doc, "核心平台高可用设计", extra_diagrams["ha"], "图：接入、服务、数据和运营四个层面的高可用机制")
        if draft.name == "ai-operations.md":
            add_figure_page(doc, "AI智能远程协助总体设计", diagrams["ai"], "图：现场感知、智能理解、协同处置与证据改进闭环")
        add_markdown(doc, draft, first_chapter=False)
        if draft.name == "ch05-08.md":
            add_figure_page(doc, "存量系统利旧与接口设计", extra_diagrams["integration"], "图：既有资源通过标准接口接入融合平台")
            add_figure_page(doc, "现网平滑迁移设计", extra_diagrams["migration"], "图：准备、试点、分批迁移和稳定移交路径")
        if draft.name == "ai-operations.md":
            add_figure_page(doc, "AI录音录像与自动报告数据流", extra_diagrams["ai-data"], "图：证据采集、智能处理、受控存储和授权使用")
            add_figure_page(doc, "统一网管与终端四自动", diagrams["automation"], "图：自动发现、自动配置、自动诊断、自动开通闭环")
            add_figure_page(doc, "全网一体化运维架构", diagrams["management"], "图：资源、监控、运维和运营中心统一管理")
            add_figure_page(doc, "纵深安全与安全审计", diagrams["security"], "图：访问、通信、数据和审计运营四域协同")

    add_scenarios(doc)
    add_tender_tables(doc, tender_doc)
    add_response_matrix(doc, requirements)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)

    text_chars = sum(len(clean_text(p.text)) for p in doc.paragraphs)
    table_chars = sum(len(clean_text(c.text)) for t in doc.tables for row in t.rows for c in row.cells)
    print(f"output={args.output}")
    print(f"paragraph_characters={text_chars}")
    print(f"table_characters={table_chars}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} requirements={len(requirements)}")


if __name__ == "__main__":
    main()
