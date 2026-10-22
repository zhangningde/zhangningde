#!/usr/bin/env python3
"""Validate a generated proposal against the tender source and rendered PDF."""

from __future__ import annotations

import argparse
import statistics
import re
import subprocess
from pathlib import Path

from docx import Document


KEYWORDS = [
    "AI降噪",
    "AI智能对话",
    "免手持",
    "安全审计",
    "录音录像",
    "自动报告",
    "统一网管",
    "自动发现",
    "自动配置",
    "自动诊断",
    "自动开通",
    "GB/T 28181",
    "RFC 3261",
    "印尼文",
]


def normalized(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def document_text(doc: Document) -> str:
    paragraphs = "\n".join(p.text for p in doc.paragraphs)
    tables = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    return normalized(paragraphs + "\n" + tables)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--proposal", type=Path, required=True)
    parser.add_argument("--tender", type=Path, required=True)
    parser.add_argument("--pdf", type=Path)
    parser.add_argument("--minimum-pages", type=int, default=320)
    parser.add_argument("--maximum-pages", type=int, default=390)
    parser.add_argument("--minimum-paragraph-characters", type=int, default=160000)
    parser.add_argument("--minimum-paragraph-table-ratio", type=float, default=1.5)
    parser.add_argument("--maximum-sparse-pages", type=int, default=25)
    parser.add_argument("--maximum-empty-pages", type=int, default=0)
    args = parser.parse_args()

    proposal = Document(args.proposal)
    tender = Document(args.tender)
    paragraph_text = normalized("\n".join(p.text for p in proposal.paragraphs))
    table_text = normalized("\n".join(cell.text for table in proposal.tables for row in table.rows for cell in row.cells))
    output = normalized(paragraph_text + "\n" + table_text)
    missing = []
    for paragraph in tender.paragraphs:
        text = normalized(paragraph.text)
        if not text or text == "目录" or paragraph.style.name.startswith(("Heading", "标题")):
            continue
        if text not in output:
            missing.append(text)

    print(f"proposal_characters={len(output)}")
    print(f"paragraph_characters={len(paragraph_text)}")
    print(f"table_characters={len(table_text)}")
    paragraph_table_ratio = len(paragraph_text) / max(len(table_text), 1)
    print(f"paragraph_to_table_ratio={paragraph_table_ratio:.2f}")
    print(f"paragraphs={len(proposal.paragraphs)} tables={len(proposal.tables)} images={len(proposal.inline_shapes)}")
    print(f"missing_tender_paragraphs={len(missing)}")
    for keyword in KEYWORDS:
        print(f"keyword[{keyword}]={output.count(keyword)}")

    if args.pdf:
        result = subprocess.run(["pdfinfo", str(args.pdf)], check=True, text=True, capture_output=True)
        pages_match = re.search(r"^Pages:\s+(\d+)", result.stdout, re.MULTILINE)
        pages = int(pages_match.group(1)) if pages_match else 0
        print(f"rendered_pages={pages}")
        if not args.minimum_pages <= pages <= args.maximum_pages:
            raise SystemExit(f"Rendered page count {pages} outside target range")
        text_result = subprocess.run(
            ["pdftotext", "-layout", str(args.pdf), "-"],
            check=True,
            text=True,
            capture_output=True,
        )
        page_texts = text_result.stdout.split("\f")
        if page_texts and not page_texts[-1].strip():
            page_texts.pop()
        densities = [len(re.sub(r"\s+", "", page)) for page in page_texts]
        empty_pages = sum(1 for value in densities if value < 20)
        sparse_pages = sum(1 for value in densities if value < 150)
        print(f"average_page_characters={statistics.mean(densities):.1f}")
        print(f"median_page_characters={statistics.median(densities):.1f}")
        print(f"sparse_pages_lt_150={sparse_pages}")
        print(f"empty_pages_lt_20={empty_pages}")
        if sparse_pages > args.maximum_sparse_pages:
            raise SystemExit(f"Sparse page count {sparse_pages} exceeds limit")
        if empty_pages > args.maximum_empty_pages:
            raise SystemExit(f"Empty page count {empty_pages} exceeds limit")

    if len(paragraph_text) < args.minimum_paragraph_characters:
        raise SystemExit(
            f"Narrative content {len(paragraph_text)} below minimum {args.minimum_paragraph_characters}"
        )
    if paragraph_table_ratio < args.minimum_paragraph_table_ratio:
        raise SystemExit(
            f"Narrative/table ratio {paragraph_table_ratio:.2f} below minimum "
            f"{args.minimum_paragraph_table_ratio:.2f}"
        )
    if missing:
        for item in missing[:20]:
            print("MISSING:", item)
        raise SystemExit("Tender paragraph coverage failed")


if __name__ == "__main__":
    main()
